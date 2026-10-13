"""Build workbooks for filling sectors with Excel's Stocks data type.

Runbook: docs/sectors-from-excel-stocks.md

    uv run python scripts/stocks_workbook.py bse Equity.csv Equity_active.xlsx
    uv run python scripts/stocks_workbook.py remaining Sectors_remaining.xlsx

``bse``: from BSE's List of Securities (Equity.csv), keeps active company shares only (drops
delisted and suspended securities, mutual fund ETFs, rights entitlements and preference shares)
and adds two columns in front: Sector (a formula) and BSE Stock ID (``XBOM:<security code>``).

``remaining``: every active company in Analekh still without a sector, with
``XNSE:<NSE symbol>`` IDs, for companies not on BSE (including NSE SME).

Convert the ID column to Stocks in Excel (Microsoft 365, desktop or web), save, then:
``uv run analekh universe sectors --stocks-file <workbook>``. The workbooks hold LSEG data
licensed to you for your own research: keep them out of git (save them outside the repo).
"""

import argparse
import contextlib
import csv
import sys
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

# The Stocks field holding the sector; FIELDVALUE needs the _xlfn prefix in the file format.
SECTOR_FORMULA = '=IFERROR(_xlfn.FIELDVALUE(B{row},"Industry"),"")'
WIDTHS = {"Sector": 30, "Issuer Name": 40, "Security Name": 34, "Company": 44, "Note": 30}


def _write(path: Path, title: str, header: list[str], rows: list[list[Any]]) -> None:
    if path.exists():
        raise SystemExit(f"{path} already exists; choose another name (nothing was overwritten)")
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append(header)
    for i, row in enumerate(rows, start=2):
        ws.append([SECTOR_FORMULA.format(row=i), *row])
    for cell in ws[1]:
        cell.font, cell.fill = Font(bold=True), PatternFill("solid", fgColor="E8F0FC")
    for j, name in enumerate(header, start=1):
        ws.column_dimensions[get_column_letter(j)].width = WIDTHS.get(name, 16)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)
    print(f"wrote {path}: {len(rows)} rows. Next: convert column B to Stocks in Excel.")


def bse(src: Path, out: Path) -> None:
    with src.open(encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))
    cols = [c for c in rows[0] if c]  # BSE adds unnamed, empty trailing columns
    for need in ("Security Code", "Status", "Instrument", "Issuer Name", "ISIN No"):
        if need not in cols:
            raise SystemExit(f"{src.name}: missing column {need!r}; found {cols}")
    keep = [
        r
        for r in rows
        if r["Status"].strip() == "Active"
        and r["Instrument"].strip() == "Equity"  # rights entitlements show "-"
        and "mutual fund" not in r["Issuer Name"].lower()  # ETFs are listed as Equity
    ]
    out_rows = []
    for r in keep:
        values: list[Any] = [(r[c] or "").strip() for c in cols]
        with contextlib.suppress(ValueError, InvalidOperation):
            values[cols.index("Face Value")] = float(Decimal(values[cols.index("Face Value")]))
        out_rows.append([f"XBOM:{r['Security Code'].strip()}", *values])
    print(f"read {len(rows)} BSE securities; kept {len(keep)} active company shares")
    _write(out, "BSE active equity", ["Sector", "BSE Stock ID", *cols], out_rows)


def remaining(out: Path) -> None:
    from sqlalchemy import select

    from analekh.db.models import Company
    from analekh.db.session import session_scope

    with session_scope() as s:
        left = s.scalars(
            select(Company)
            .where(Company.is_active, Company.sector.is_(None))
            .order_by(Company.is_sme, Company.nse_symbol)
        ).all()
        rows = [
            [
                f"XNSE:{c.nse_symbol}",
                c.nse_symbol,
                c.name,
                c.isin,
                c.series,
                "SME" if c.is_sme else "Main board",
                "NSE SME: not on BSE" if c.is_sme else "",
            ]
            for c in left
        ]
    header = ["Sector", "NSE Stock ID", "NSE Symbol", "Company", "ISIN", "Series", "Board", "Note"]
    _write(out, "Sectors remaining", header, rows)


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    sub = ap.add_subparsers(dest="cmd", required=True)
    b = sub.add_parser("bse", help="BSE List of Securities -> workbook with XBOM: IDs")
    b.add_argument("equity_csv", type=Path)
    b.add_argument("out", type=Path)
    r = sub.add_parser("remaining", help="companies without a sector -> workbook with XNSE: IDs")
    r.add_argument("out", type=Path)
    args = ap.parse_args()
    if args.cmd == "bse":
        bse(args.equity_csv, args.out)
    else:
        remaining(args.out)


if __name__ == "__main__":
    main()
