"""Build formatted design document (.docx) from docs/design/technical-design.md, using previous
version's Word file as template: styles, list numbering, title block, table look, landscape
section around Figure 1. Figures = diagram panels of architecture.html (scripts/diagram_panels.py
writes panel1..N.png); each gets a landscape page, heading just above when there is one. Numbered
lists ("1. ...") restart at 1. Headers and footers get new version number.

    uv run python scripts/diagram_panels.py file:///.../docs/design/architecture.html out/figs
    uv run python scripts/design_docx.py "docs/design/<previous>.docx"
        docs/design/technical-design.md "docs/design/<new>.docx" out/figs   (one command)

Bump OLD and NEW below (and template argument) for next version.
"""

import copy
import re
import sys

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu

TEMPLATE, MARKDOWN, OUT, FIGS = sys.argv[1:5]
OLD, NEW = "0.0.5", "0.0.6"  # the template's version, this document's

d = docx.Document(TEMPLATE)
body = d.element.body
kids = list(body)


def sect_of(child):
    return copy.deepcopy(child.find(".//" + qn("w:sectPr")))


breaks = [k for k in kids if k.tag == qn("w:p") and k.find(".//" + qn("w:sectPr")) is not None]
portrait_break, landscape_break = sect_of(breaks[0]), sect_of(breaks[1])
title_block = [copy.deepcopy(k) for k in kids[:3]]  # Analekh / tagline / Technical Design Document
first_style_ppr = {}
for k in kids:
    if k.tag != qn("w:p"):
        continue
    ps = k.find(qn("w:pPr") + "/" + qn("w:pStyle"))
    if ps is not None and ps.get(qn("w:val")) not in first_style_ppr:
        first_style_ppr[ps.get(qn("w:val"))] = copy.deepcopy(k.find(qn("w:pPr")))
final_sect = body.find(qn("w:sectPr"))
for k in kids:
    if k is not final_sect:
        body.remove(k)

# Numbered lists in v0.0.2, by the start of their first item; other lists are bullets (numId 3).
NUMBERED = {
    "Search ticker": "4",
    "Start coverage on any company": "5",
    "Upload a workbook": "6",
    "Local read": "7",
    "Hidden tag": "8",
    "Pick a company from the universe": "9",
    "Clone the repository": "10",
}


def add(el):
    final_sect.addprevious(el)
    return el


def ppr(spacing=None, extra=()):
    p = OxmlElement("w:p")
    pp = OxmlElement("w:pPr")
    for e in extra:
        pp.append(e)
    if spacing:
        sp = OxmlElement("w:spacing")
        for k, v in spacing.items():
            sp.set(qn(f"w:{k}"), str(v))
        pp.append(sp)
    p.append(pp)
    return p


def rpr(bold=False, italic=False, size=None, color=None, font=None, underline=False):
    r = OxmlElement("w:rPr")
    if font:
        f = OxmlElement("w:rFonts")
        for a in ("ascii", "hAnsi", "cs", "eastAsia"):
            f.set(qn(f"w:{a}"), font)
        r.append(f)
    if bold:
        r.append(OxmlElement("w:b"))
        r.append(OxmlElement("w:bCs"))
    if italic:
        r.append(OxmlElement("w:i"))
        r.append(OxmlElement("w:iCs"))
    if color:
        c = OxmlElement("w:color")
        c.set(qn("w:val"), color)
        r.append(c)
    if size:
        for tag in ("w:sz", "w:szCs"):
            s = OxmlElement(tag)
            s.set(qn("w:val"), str(size))
            r.append(s)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        r.append(u)
    return r


def run(text, **fmt):
    r = OxmlElement("w:r")
    props = rpr(**fmt)
    if len(props):
        r.append(props)
    t = OxmlElement("w:t")
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    r.append(t)
    return r


TOKEN = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|(?<![\w*])\*[^*\s][^*]*?\*(?![\w*]))")


def inline(p, text, size=None, italic=False, color=None, bold=False):
    """Markdown bold, italic, code and links as runs."""
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            p.append(run(part[2:-2], bold=True, size=size, italic=italic, color=color))
        elif part.startswith("`") and part.endswith("`"):
            p.append(run(part[1:-1], font="Consolas", size=(size or 21) - 2, color=color))
        elif part.startswith("[") and "](" in part:
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part).groups()
            rid = d.part.relate_to(url, RT.HYPERLINK, is_external=True)
            h = OxmlElement("w:hyperlink")
            h.set(qn("r:id"), rid)
            h.append(run(label, color="1F5FBF", underline=True, size=size))
            p.append(h)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            p.append(run(part[1:-1], italic=True, size=size, bold=bold, color=color))
        else:
            p.append(run(part, size=size, italic=italic, color=color, bold=bold))


def paragraph(text):
    p = ppr({"after": 140})
    inline(p, text)
    add(p)


def heading(level, text, page_break=True):
    style = {1: "Heading1", 2: "Heading2", 3: "Heading3"}[level]
    p = OxmlElement("w:p")
    pp = copy.deepcopy(first_style_ppr.get(style)) if style in first_style_ppr else OxmlElement("w:pPr")
    if not page_break:
        for e in pp.findall(qn("w:pageBreakBefore")):
            pp.remove(e)
    p.append(pp)
    p.append(run(text))
    add(p)


numbering = d.part.numbering_part.element
_base = next(n for n in numbering.findall(qn("w:num")) if n.get(qn("w:numId")) == "4")
NUMBERED_ABSTRACT = _base.find(qn("w:abstractNumId")).get(qn("w:val"))


def new_numbered_list():
    """A fresh numbering instance of the numbered-list style, starting again at 1."""
    nid = str(max(int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))) + 1)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), nid)
    ab = OxmlElement("w:abstractNumId")
    ab.set(qn("w:val"), NUMBERED_ABSTRACT)
    num.append(ab)
    lo = OxmlElement("w:lvlOverride")
    lo.set(qn("w:ilvl"), "0")
    so = OxmlElement("w:startOverride")
    so.set(qn("w:val"), "1")
    lo.append(so)
    num.append(lo)
    numbering.append(num)
    return nid


def list_item(text, num_id):
    p = OxmlElement("w:p")
    pp = OxmlElement("w:pPr")
    ps = OxmlElement("w:pStyle")
    ps.set(qn("w:val"), "ListParagraph")
    pp.append(ps)
    np_ = OxmlElement("w:numPr")
    il = OxmlElement("w:ilvl")
    il.set(qn("w:val"), "0")
    ni = OxmlElement("w:numId")
    ni.set(qn("w:val"), num_id)
    np_.append(il)
    np_.append(ni)
    pp.append(np_)
    sp = OxmlElement("w:spacing")
    sp.set(qn("w:after"), "60")
    pp.append(sp)
    p.append(pp)
    inline(p, text)
    add(p)


def code_block(lines):
    for i, line in enumerate(lines):
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "F3F5F2")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:val"), "clear")
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "120")
        ind.set(qn("w:right"), "120")
        p = ppr(
            {
                "after": 120 if i == len(lines) - 1 else 0,
                "before": 80 if i == 0 else 0,
                "line": 240,
            },
            extra=[shd],
        )
        p.find(qn("w:pPr")).append(ind)
        p.append(run(line, font="Consolas", size=16))
        add(p)


def cell_border(tcpr):
    b = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:color"), "C7D0CA")
        e.set(qn("w:sz"), "4")
        b.append(e)
    tcpr.append(b)


def table(rows, width=9638):
    header, body_rows = rows[0], rows[1:]
    ncol = len(header)
    weights = [max(6, min(60, max(len(re.sub(r"[*`]", "", r[i] if i < len(r) else "")) for r in rows))) for i in range(ncol)]
    total = sum(weights)
    widths = [max(900, int(width * w / total)) for w in weights]
    widths[-1] += width - sum(widths)
    tbl = OxmlElement("w:tbl")
    tp = OxmlElement("w:tblPr")
    tw = OxmlElement("w:tblW")
    tw.set(qn("w:type"), "dxa")
    tw.set(qn("w:w"), str(width))
    tp.append(tw)
    tb = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:color"), "auto")
        e.set(qn("w:sz"), "4")
        tb.append(e)
    tp.append(tb)
    tbl.append(tp)
    grid = OxmlElement("w:tblGrid")
    for w in widths:
        g = OxmlElement("w:gridCol")
        g.set(qn("w:w"), str(w))
        grid.append(g)
    tbl.append(grid)
    for ri, row in enumerate([header] + body_rows):
        tr = OxmlElement("w:tr")
        if ri == 0:
            trp = OxmlElement("w:trPr")
            trp.append(OxmlElement("w:cantSplit"))
            trp.append(OxmlElement("w:tblHeader"))
            tr.append(trp)
        for ci in range(ncol):
            text = row[ci] if ci < len(row) else ""
            tc = OxmlElement("w:tc")
            tcp = OxmlElement("w:tcPr")
            w_ = OxmlElement("w:tcW")
            w_.set(qn("w:type"), "dxa")
            w_.set(qn("w:w"), str(widths[ci]))
            tcp.append(w_)
            cell_border(tcp)
            if ri == 0:
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), "E6EDE9")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:val"), "clear")
                tcp.append(shd)
            mar = OxmlElement("w:tcMar")
            for side, v in (("top", 60), ("left", 100), ("bottom", 60), ("right", 100)):
                e = OxmlElement(f"w:{side}")
                e.set(qn("w:type"), "dxa")
                e.set(qn("w:w"), str(v))
                mar.append(e)
            tcp.append(mar)
            tc.append(tcp)
            p = ppr({"after": 0})
            clean = text.replace("**", "") if ri == 0 else text
            inline(p, clean, size=19, bold=ri == 0)
            tc.append(p)
            tr.append(tc)
        tbl.append(tr)
    add(tbl)
    add(ppr({"after": 80}))  # space after the table, as in the template


def section_break(sect):
    p = OxmlElement("w:p")
    pp = OxmlElement("w:pPr")
    pp.append(copy.deepcopy(sect))
    p.append(pp)
    add(p)


def figure(path, width_emu):
    from PIL import Image

    w, h = Image.open(path).size
    p = d.add_paragraph()  # python-docx builds the picture XML and relationship
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Emu(76200)
    p.paragraph_format.space_after = Emu(50800)
    p.add_run().add_picture(path, width=Emu(width_emu), height=Emu(int(width_emu * h / w)))
    el = p._p
    el.getparent().remove(el)
    add(el)


def caption(text):
    p = ppr({"after": 200})
    inline(p, text.strip("*"), size=18, italic=True, color="5A6862")
    add(p)


def cells(line):
    line = line.strip().strip("|")
    parts = re.split(r"(?<!\\)\|", line)
    return [c.strip().replace("\\|", "|") for c in parts]


with open(MARKDOWN, encoding="utf-8") as fh:
    lines = fh.read().splitlines()
for el in title_block:
    add(el)
i = 0
# after "Technical Design Document".
while not lines[i].startswith("Technical Design Document"):
    i += 1
i += 1
figure_n = 0
in_contents = False
after_break = False
landscape_open = False  # a landscape section was opened for the next figure


def next_line(j):
    while j < len(lines) and not lines[j].strip():
        j += 1
    return lines[j].strip() if j < len(lines) else ""


while i < len(lines):
    line = lines[i]
    s = line.strip()
    if not s:
        i += 1
        continue
    if s.startswith("```"):
        block = []
        i += 1
        while i < len(lines) and not lines[i].strip().startswith("```"):
            block.append(lines[i])
            i += 1
        code_block(block or [""])
        i += 1
        continue
    if s.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            if not re.match(r"^\|\s*:?-{3,}", lines[i].strip()):
                rows.append(cells(lines[i]))
            i += 1
        table(rows)
        continue
    m = re.match(r"^(#{1,4})\s+(.*)$", s)
    if m:
        hashes, text = m.groups()
        level = {1: 1, 2: 1, 3: 2, 4: 3}[len(hashes)]
        if text.startswith("5. System architecture"):
            section_break(portrait_break)  # the architecture chapter opens a landscape section
            after_break = True
        elif next_line(i + 1).startswith("*Figure"):
            section_break(portrait_break)  # the heading goes on its figure's landscape page
            after_break = landscape_open = True
        heading(level, text, page_break=not after_break)
        after_break = False
        in_contents = text == "Contents"
        i += 1
        continue
    if re.match(r"^\d+\. ", s):
        num_id = new_numbered_list()
        while i < len(lines) and re.match(r"^\d+\. ", lines[i].strip()):
            list_item(re.sub(r"^\d+\.\s+", "", lines[i].strip()), num_id)
            i += 1
        continue
    if s.startswith("- "):
        first = re.sub(r"[*`]", "", s[2:])
        num_id = "2" if in_contents else next((v for k, v in NUMBERED.items() if first.startswith(k)), "3")
        while i < len(lines) and lines[i].strip().startswith("- "):
            list_item(lines[i].strip()[2:], num_id)
            i += 1
        continue
    if s.startswith("*Figure"):
        figure_n += 1
        if figure_n == 1:
            figure(f"{FIGS}/panel1.png", 8591550)
            caption(s)
            section_break(landscape_break)
            after_break = True
        else:
            if not landscape_open:
                section_break(portrait_break)
            figure(f"{FIGS}/panel{figure_n}.png", 8591550)
            caption(s)
            section_break(landscape_break)
            after_break, landscape_open = True, False
        i += 1
        continue
    para = [s]
    i += 1
    while i < len(lines) and lines[i].strip() and not re.match(r"^(#|\||- |```|\*Figure|\d+\. )", lines[i].strip()):
        para.append(lines[i].strip())
        i += 1
    paragraph(" ".join(para))

for sec in d.sections:
    for part in (sec.footer, sec.header, sec.first_page_footer, sec.first_page_header):
        for t in part._element.iter(qn("w:t")):
            if t.text and OLD in t.text:
                t.text = t.text.replace(OLD, NEW)

cp = d.core_properties
cp.title = "Analekh Technical Design Document"
cp.subject = f"Version {NEW}"
cp.author = "Kaushal Attal"
d.save(OUT)
print("saved", OUT)
