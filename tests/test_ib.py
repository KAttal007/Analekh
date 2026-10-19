import io
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook
from pptx import Presentation
from pptx.util import Inches

from analekh.db.models import Company, Price
from analekh.ib.comps import CompInputs, compute, gather, summary_stats, to_xlsx
from analekh.ib.pitch import refresh
from analekh.ib.profile import ProfileData, build
from analekh.ingest.pipeline import ingest_results_xbrl
from tests.xbrl_factory import build_instance, pnl

CR = Decimal(10_000_000)


def comp(symbol, mcap, pat, revenue=None, debt=None, cash=None, book=None):
    return CompInputs(
        symbol,
        symbol.title(),
        mcap * CR,
        debt and debt * CR,
        cash and cash * CR,
        revenue and revenue * CR,
        None,
        pat and pat * CR,
        book and book * CR,
    )


def test_multiples_and_not_meaningful():
    row = compute(comp("AAA", 10_000, 500, revenue=4_000, debt=2_000, cash=1_000, book=5_000))
    assert row.ev == 11_000 * CR
    assert row.pe == Decimal("20.0") and row.pb == Decimal("2.0")
    assert row.ev_sales == Decimal("2.8") and row.ev_ebitda is None
    loss = compute(comp("BBB", 5_000, -100))
    assert loss.pe is None  # negative earnings: n.m., not a negative multiple
    # Without debt and cash there is no EV: market cap is never passed off as EV.
    unsourced = compute(comp("CCC", 10_000, 500, revenue=4_000))
    assert unsourced.ev is None and unsourced.ev_sales is None


def test_summary_and_xlsx_export():
    rows = [
        compute(comp("A", 1000, 50)),
        compute(comp("B", 3000, 100)),
        compute(comp("C", 2000, -5)),
    ]
    assert summary_stats(rows)["pe"]["median"] == Decimal("25.0")
    ws = load_workbook(io.BytesIO(to_xlsx(rows)))["Comps"]
    assert ws["K4"].value == "P/E"
    assert ws["A5"].value == "A" and ws["K5"].value == 20.0 and ws["K7"].value == "n.m."
    assert ws["B4"].value == "Basis" and ws["B5"].value == "n.a."  # inputs built by hand


def test_pitch_placeholders_replaced_and_logged():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tf = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text_frame
    para = tf.paragraphs[0]
    for part in ("PAT of {{DEMO.", "pat_q}} crore; ", "{{DEMO.unknown}}"):
        para.add_run().text = part
    buf = io.BytesIO()
    prs.save(buf)

    out, log = refresh(buf.getvalue(), lambda s, m: "187.5" if m == "pat_q" else None)
    text = Presentation(io.BytesIO(out)).slides[0].shapes[0].text_frame.text
    assert text == "PAT of 187.5 crore; {{DEMO.unknown}}"
    assert [(r.placeholder, r.value) for r in log] == [
        ("{{DEMO.pat_q}}", "187.5"),
        ("{{DEMO.unknown}}", None),
    ]


def test_profile_slide_builds():
    data = ProfileData(
        "Demo Finance",
        "DEMOFIN",
        "Gold loan NBFC.",
        [("Revenue", ["900", "1,000"]), ("PAT", ["165", "188"])],
        ["Q1FY27", "Q2FY27"],
        [("Market cap", "n.a.")],
        ["Board meeting"],
    )
    prs = Presentation(io.BytesIO(build(data)))
    assert len(prs.slides) == 1


QUARTER_ENDS = [date(2025, 9, 30), date(2025, 12, 31), date(2026, 3, 31), date(2026, 6, 30)]


def _file_quarters(db, store, company, nature, revenue, ends=QUARTER_ENDS):
    for end in ends:
        start = date(end.year, end.month - 2, 1)
        q = pnl(revenue * 10_000_000, 0, revenue * 8_000_000, revenue * 500_000)
        data = build_instance(q_start=start, q_end=end, ytd_start=start, values_q=q, nature=nature)
        ingest_results_xbrl(db, store, company, data)


def _company(db, symbol):
    c = Company(
        nse_symbol=symbol,
        name=f"{symbol.title()} Limited",
        isin=f"INE{symbol[:4]}01011",  # 12 characters, like a real ISIN
    )
    db.add(c)
    db.flush()
    return c


def test_gather_falls_back_to_standalone_only_filers(db, store):
    solo = _company(db, "SOLOCO")
    _file_quarters(db, store, solo, "Standalone", 100)
    inputs = gather(db, "SOLOCO")
    assert inputs.consolidated is False and inputs.ttm_revenue == 400 * CR
    assert ws_basis(inputs) == "Standalone"


def test_gather_prefers_consolidated_and_can_be_fixed(db, store):
    both = _company(db, "BOTHCO")
    _file_quarters(db, store, both, "Consolidated", 120)
    _file_quarters(db, store, both, "Standalone", 100)
    assert gather(db, "BOTHCO").consolidated is True
    assert gather(db, "BOTHCO").ttm_revenue == 480 * CR
    fixed = gather(db, "BOTHCO", consolidated=False)
    assert fixed.consolidated is False and fixed.ttm_revenue == 400 * CR


def test_gather_uses_standalone_when_consolidated_is_short(db, store):
    # Like ABB: consolidated for only the latest two quarters, standalone for all four.
    switch = _company(db, "SWITCHCO")
    _file_quarters(db, store, switch, "Consolidated", 120, QUARTER_ENDS[2:])
    _file_quarters(db, store, switch, "Standalone", 100)
    inputs = gather(db, "SWITCHCO")
    assert inputs.consolidated is False and inputs.ttm_revenue == 400 * CR
    assert "standalone: fewer than four consolidated quarters" in inputs.notes


def ws_basis(inputs):
    ws = load_workbook(io.BytesIO(to_xlsx([compute(inputs)])))["Comps"]
    return ws["B5"].value


def test_ttm_needs_four_back_to_back_quarters(db, store):
    gap = _company(db, "GAPCO")  # December 2025 never filed
    ends = [date(2025, 6, 30), date(2025, 9, 30), date(2026, 3, 31), date(2026, 6, 30)]
    _file_quarters(db, store, gap, "Standalone", 100, ends)
    inputs = gather(db, "GAPCO")
    assert inputs.ttm_revenue is None and inputs.consolidated is False  # still says which basis
    assert "fewer than four back-to-back quarters of results on record" in inputs.notes


def test_comps_row_has_margins_growth_book_value_and_no_ebitda_for_lenders(db, store):
    co = _company(db, "FULLCO")
    ends = [date(2025, 6, 30), *QUARTER_ENDS]  # five quarters, so June has a year-ago quarter
    for i, end in enumerate(ends):
        revenue = (100 + 10 * i) * 10_000_000
        q = pnl(revenue, 0, revenue * 8 // 10, revenue // 20)  # PBT 20%, PAT 15% of revenue
        q |= {
            "FinanceCosts": revenue // 100,
            "DepreciationDepletionAndAmortisationExpense": revenue // 50,
        }
        last = end == ends[-1]
        data = build_instance(
            q_start=date(end.year, end.month - 2, 1),
            q_end=end,
            ytd_start=date(end.year, end.month - 2, 1),
            values_q=q,
            values_instant={"Equity": 1_000 * 10_000_000} if last else None,
            nature="Standalone",
        )
        ingest_results_xbrl(db, store, co, data)
    db.add(
        Price(
            company_id=co.id,
            trade_date=date(2026, 10, 16),
            close=Decimal(500),
            market_cap=2_000 * CR,
        )
    )
    db.flush()

    row = compute(gather(db, "FULLCO"))
    i = row.inputs
    assert i.ttm_revenue == 500 * CR  # 110 + 120 + 130 + 140
    assert i.ebitda_margin == Decimal("23.0")  # PBT 20% + finance 1% + D&A 2%
    assert i.pat_margin == Decimal("15.0") and i.revenue_yoy == Decimal("40.0")
    assert i.book_value == 1_000 * CR and i.book_value_at == date(2026, 6, 30)
    assert row.pb == Decimal("2.0") and row.pe == Decimal("26.7")  # 2,000 / 75
    assert row.ev is None and row.ev_ebitda is None

    co.sector_lseg = "Banking Services"  # a lender: EBITDA means nothing
    db.flush()
    assert gather(db, "FULLCO").ttm_ebitda is None
