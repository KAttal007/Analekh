"""Excel bridge: linking must not rewrite the analyst's model, and open files must not vanish."""

import uuid
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.packaging.custom import StringProperty

from analekh.excel_bridge import indexer, writer
from analekh.workbook.tags import (
    Tag,
    WorkbookLocked,
    read_tag,
    write_tag_file,
)

TAG = Tag(uuid.UUID("11111111-1111-1111-1111-111111111111"), "LEMONTREE")


def _model(path: Path, custom: bool = False) -> Path:
    wb = Workbook()
    wb.active["A1"] = "Lemon Tree Hotels"
    wb.active["B2"] = 5
    if custom:
        wb.custom_doc_props.append(StringProperty(name="analyst_note", value="keep me"))
    wb.save(path)
    return path


def _parts(path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n) for n in z.namelist()}


@pytest.mark.parametrize("custom", [False, True])
def test_tag_changes_only_the_tag_parts(tmp_path, custom):
    path = _model(tmp_path / "m.xlsx", custom)
    before = _parts(path)
    write_tag_file(path, TAG)
    after = _parts(path)
    changed = {n for n in after if before.get(n) != after[n]}
    assert changed <= {"docProps/custom.xml", "[Content_Types].xml", "_rels/.rels"}
    assert set(before) <= set(after)
    wb = load_workbook(path)
    assert read_tag(wb) == TAG
    assert wb.active["B2"].value == 5
    if custom:
        assert "analyst_note" in wb.custom_doc_props.names


def test_tag_is_replaced_not_duplicated(tmp_path):
    path = _model(tmp_path / "m.xlsx")
    write_tag_file(path, TAG)
    other = Tag(uuid.uuid4(), "TCS")
    write_tag_file(path, other)
    wb = load_workbook(path)
    assert read_tag(wb) == other
    assert len(wb.custom_doc_props.props) == 2
    assert not list(tmp_path.glob(".*analekh-tmp*"))


def test_locked_file_is_left_untouched(tmp_path, monkeypatch):
    path = _model(tmp_path / "m.xlsx")
    before = path.read_bytes()

    def locked(*_a, **_k):
        raise PermissionError("sharing violation")

    monkeypatch.setattr("analekh.workbook.tags.os.replace", locked)
    with pytest.raises(WorkbookLocked):
        write_tag_file(path, TAG)
    assert path.read_bytes() == before
    assert not list(tmp_path.glob(".*analekh-tmp*"))


def test_link_tag_says_so_when_excel_holds_the_file(tmp_path, monkeypatch):
    path = _model(tmp_path / "m.xlsx")
    monkeypatch.setattr("analekh.workbook.tags.os.replace", _raise_locked)
    monkeypatch.setattr(writer, "_tag_open_book", lambda *_a: False)
    with pytest.raises(RuntimeError, match="open in another program"):
        writer.link_tag(path, TAG)
    monkeypatch.setattr(writer, "_tag_open_book", lambda *_a: True)
    assert writer.link_tag(path, TAG) == "open"


def _raise_locked(*_a, **_k):
    raise PermissionError("sharing violation")


def test_scan_reads_a_copy_when_the_file_is_locked(tmp_path, monkeypatch):
    path = _model(tmp_path / "m.xlsx")
    write_tag_file(path, TAG)
    real = indexer._load

    def load(p):
        if Path(p) == path:
            raise PermissionError("locked by Excel")
        return real(p)

    monkeypatch.setattr(indexer, "_load", load)
    skipped: list[tuple[Path, str]] = []
    [hint] = indexer.scan([tmp_path], skipped)
    assert hint.tag_symbol == "LEMONTREE"
    assert hint.path == str(path)
    assert not skipped


def test_scan_reports_what_it_cannot_read(tmp_path):
    bad = tmp_path / "broken.xlsx"
    bad.write_bytes(b"not a workbook")
    skipped: list[tuple[Path, str]] = []
    assert indexer.scan([tmp_path], skipped) == []
    [(path, reason)] = skipped
    assert path == bad and reason


def test_scan_ignores_excel_lock_and_temp_files(tmp_path):
    _model(tmp_path / "~$m.xlsx")
    _model(tmp_path / ".m.analekh-tmp.xlsx")
    assert indexer.scan([tmp_path]) == []


def test_scan_asks_excel_for_a_copy_when_even_copying_is_refused(tmp_path, monkeypatch):
    path = _model(tmp_path / "m.xlsx")
    write_tag_file(path, TAG)
    saved = path.read_bytes()
    real = indexer._load

    def load(p):
        if Path(p) == path:
            raise PermissionError("locked by Excel")
        return real(p)

    def refuse(*_a, **_k):
        raise PermissionError("locked by Excel")

    def excel_copy(src, dest):
        assert src == path
        dest.write_bytes(saved)
        return True

    monkeypatch.setattr(indexer, "_load", load)
    monkeypatch.setattr(indexer.shutil, "copyfile", refuse)
    monkeypatch.setattr(indexer.live, "save_copy", excel_copy)
    [hint] = indexer.scan([tmp_path])
    assert hint.tag_symbol == "LEMONTREE"


def test_scan_reports_a_locked_file_excel_cannot_copy(tmp_path, monkeypatch):
    path = _model(tmp_path / "m.xlsx")

    def refuse(*_a, **_k):
        raise PermissionError("locked")

    monkeypatch.setattr(indexer, "_load", refuse)
    monkeypatch.setattr(indexer.live, "save_copy", lambda *_a: False)
    skipped: list[tuple[Path, str]] = []
    assert indexer.scan([tmp_path], skipped) == []
    assert skipped[0][0] == path and "PermissionError" in skipped[0][1]


def test_cli_says_plainly_when_the_token_is_missing():
    import httpx
    import typer

    from analekh.excel_bridge.cli import _ok

    req = httpx.Request("POST", "http://127.0.0.1:8000/models/discover")
    with pytest.raises(typer.Exit):
        _ok(httpx.Response(401, request=req))
    with pytest.raises(typer.Exit):
        _ok(httpx.Response(404, json={"detail": "unknown symbol"}, request=req))
    assert _ok(httpx.Response(200, json={}, request=req)).status_code == 200


def test_a_file_excel_will_not_even_let_us_read_counts_as_locked(tmp_path, monkeypatch):
    import zipfile as zf

    path = _model(tmp_path / "m.xlsx")

    def refuse(*_a, **_k):
        raise PermissionError("locked by Excel")

    monkeypatch.setattr("analekh.workbook.tags.zipfile.ZipFile", refuse)
    with pytest.raises(WorkbookLocked):
        write_tag_file(path, TAG)
    assert zf.is_zipfile(path)
    monkeypatch.setattr(writer, "_tag_open_book", lambda *_a: True)
    assert writer.link_tag(path, TAG) == "open"
