"""API flows over HTTP (design doc 20) on PostgreSQL with a local object store. SYNTHETIC data."""

import io
from datetime import date

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import delete, select

from analekh.api.deps import get_current, get_store, get_user
from analekh.api.main import app
from analekh.db.models import Company, Filing, Model, PeerList, Price, Watchlist
from analekh.db.models import User as UserRow
from analekh.db.session import get_session
from analekh.facts import DbValueSource
from analekh.ingest.pipeline import ingest_results_xbrl
from analekh.modelsync.changeset import build_change_set, save_proposal
from analekh.templates_builtin import NBFC
from analekh.templates_builtin import build as build_builtin
from analekh.universe.loader import parse_equity_list, upsert_universe
from analekh.workbook.mapping import Mapping
from tests.test_builder import QUARTERS, history
from tests.test_universe import EQUITY_L
from tests.xbrl_factory import BS, Q2, build_instance, pnl, quarter_instance

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture
def client(db, store):
    app.dependency_overrides[get_session] = lambda: db
    app.dependency_overrides[get_store] = lambda: store
    # Sign-in is tested in test_auth_api.py; here every request is a fixed admin's.
    admin = UserRow(id="local", email="admin@test.local", name="Admin", role="admin")
    if db.get(UserRow, "local") is None:  # a saved user: model locks refer to it
        db.add(admin)
        db.flush()
    app.dependency_overrides[get_current] = lambda: admin
    app.dependency_overrides[get_user] = lambda: "local"
    yield TestClient(app)
    app.dependency_overrides.clear()


@pytest.fixture
def demo(db, store):
    upsert_universe(db, parse_equity_list(EQUITY_L))
    c = db.scalar(select(Company).where(Company.nse_symbol == "DEMOFIN"))
    c.sector, c.industry = "Financial Services", "NBFC"
    for i, end in enumerate(QUARTERS):
        ingest_results_xbrl(db, store, c, quarter_instance(end, **history(i)))
    return c


def test_company_index_for_in_browser_search(client, demo):
    first = client.get("/companies/index")
    rows = {r["nse_symbol"]: r for r in first.json()}
    assert rows["DEMOFIN"] == {
        "id": demo.id,
        "nse_symbol": "DEMOFIN",
        "name": "Demo Finance Limited",
        "isin": "INE000D01011",
        "sector": "Financial Services",
    }
    etag = first.headers["ETag"]
    again = client.get("/companies/index", headers={"If-None-Match": etag})
    assert again.status_code == 304 and again.headers["ETag"] == etag


def test_search_company_page_and_watchlist(client, demo, db):
    db.execute(delete(Watchlist))  # the owner's own watchlist, in a rolled-back transaction
    hits = client.get("/companies/search", params={"q": "demo fin"}).json()
    assert hits[0]["nse_symbol"] == "DEMOFIN"
    page = client.get("/companies/demofin").json()
    assert page["coverage_level"] == "numbers" and not page["watched"]
    assert client.get("/companies/OLDCO").json()["coverage_level"] == "universe"
    assert client.get("/companies/NOPE").status_code == 404

    added = client.post("/watchlist", json={"symbol": "DEMOFIN", "settings": {"alerts": True}})
    assert added.json()["coverage_level"] == "deep"
    [entry] = client.get("/watchlist").json()
    assert entry["symbol"] == "DEMOFIN" and entry["last_filing"] == "2026-09-30"
    assert client.delete("/watchlist/DEMOFIN").status_code == 204
    assert client.get("/watchlist").json() == []

    filings = client.get("/companies/DEMOFIN/filings").json()
    assert filings[0]["period_end"] == "2026-09-30" and len(filings) == 8
    facts = client.get(
        "/companies/DEMOFIN/facts",
        params={"elements": "RevenueFromOperations", "start": "2026-09-30"},
    ).json()
    assert [f["value"] for f in facts] == ["975000000.0000"]


def test_facts_fall_back_to_standalone_when_company_files_no_consolidated(client, db, store):
    solo = Company(nse_symbol="SOLO", name="Solo Systems Limited", isin="INE000S01011")
    db.add(solo)
    db.flush()
    q = pnl(500_000_000, 10_000_000, 400_000_000, 25_000_000)
    standalone = build_instance(**Q2, values_q=q, values_instant=BS, nature="Standalone")
    ingest_results_xbrl(db, store, solo, standalone)
    facts = client.get("/companies/SOLO/facts", params={"elements": "RevenueFromOperations"})
    assert {f["value"] for f in facts.json()} == {"500000000.0000"}  # quarter and YTD
    asked = client.get(
        "/companies/SOLO/facts", params={"elements": "RevenueFromOperations", "consolidated": True}
    )
    assert asked.json() == []


def test_start_coverage_assumptions_finalize_then_review_and_apply(client, db, store, demo):
    built = client.post("/coverage", json={"symbol": "DEMOFIN"}).json()
    assert built["build_status"] == "awaiting_assumptions" and built["filled_cells"] >= 32
    model_id = built["model"]["id"]
    assert client.post(f"/models/{model_id}/finalize").status_code == 409  # pending assumptions

    tax = next(a for a in built["assumptions"] if a["driver"] == "tax_rate")
    decided = client.post(
        f"/models/{model_id}/assumptions/decisions",
        json={
            "decisions": [{"id": tax["id"], "decision": "edited", "final_value": "0.25"}],
            "accept_all": True,
        },
    ).json()
    assert {a["decision"] for a in decided} == {"accepted", "edited"}
    done = client.post(f"/models/{model_id}/finalize").json()
    assert done["status"] == "built" and done["checks"]["passed"] and done["version"] == 1

    # Next results: change set on the built model, reviewed and applied as a new version.
    nxt = ingest_results_xbrl(db, store, demo, quarter_instance(date(2026, 12, 31), **history(8)))
    model = db.get(Model, model_id)
    v1 = client.get(f"/models/{model_id}/versions/1")
    proposal = build_change_set(
        load_workbook(io.BytesIO(v1.content)),
        Mapping.model_validate(model.mapping),
        date(2026, 12, 31),
        DbValueSource(db, demo.id),
        last_fingerprint=model.fingerprint,
    )
    cs = save_proposal(db, model_id, 1, nxt.filing.id, proposal)

    assert client.post(f"/change-sets/{cs.id}/apply").status_code == 409  # nothing decided
    items = client.get(f"/change-sets/{cs.id}").json()["items"]
    revenue = next(i for i in items if i["kind"] == "actual_write" and i["cell"] == "J3")
    r = client.post(
        f"/change-sets/{cs.id}/decisions",
        json={
            "accept_kinds": ["actual_write", "roll_forward", "label_shift"],
            "decisions": [{"item_id": revenue["id"], "decision": "edited", "final_value": "100.5"}],
        },
    )
    remaining = [i for i in r.json()["items"] if i["decision"] == "pending"]
    if remaining:
        r = client.post(
            f"/change-sets/{cs.id}/decisions",
            json={"decisions": [{"item_id": i["id"], "decision": "skipped"} for i in remaining]},
        )
    assert r.json()["status"] == "approved"
    applied = client.post(f"/change-sets/{cs.id}/apply").json()
    assert applied["version"] == 2
    ws = load_workbook(io.BytesIO(client.get(f"/models/{model_id}/versions/2").content))["Data"]
    assert ws["J3"].value == 100.5 and ws["J2"].value == "A"
    v1_ws = load_workbook(io.BytesIO(v1.content))["Data"]
    assert isinstance(v1_ws["J3"].value, str)  # version 1 untouched


def test_template_upload_confirm_test_activate(client, demo):
    data, _ = build_builtin(NBFC, date(2026, 9, 30))
    t = client.post(
        "/templates",
        data={"name": "house_nbfc", "sectors": "NBFC"},
        files={"workbook": ("house.xlsx", data, XLSX)},
    ).json()
    assert t["status"] == "draft" and t["mapping"]["rows"]
    assert client.post(f"/templates/{t['id']}/activate").status_code == 409
    mapping = t["mapping"] | {
        "checks": [
            {"sheet": "Checks", "cell_label": "All checks pass", "expect": 1, "tolerance": 0}
        ]
    }
    assert client.post(f"/templates/{t['id']}/mapping", json=mapping).status_code == 200
    test = client.post(f"/templates/{t['id']}/test", json={"symbol": "DEMOFIN"}).json()
    assert test["result"] == "pass", test["issues"]
    assert client.post(f"/templates/{t['id']}/activate").json()["status"] == "active"
    names = {x["name"] for x in client.get("/templates", params={"sector": "NBFC"}).json()}
    assert {"house_nbfc", "nbfc_quarterly"} <= names
    bad = client.post(
        "/templates",
        data={"name": "x", "sectors": "NBFC"},
        files={"workbook": ("x.xlsx", b"not a workbook", XLSX)},
    )
    assert bad.status_code == 422


def test_discover_and_link(client, db, demo):
    found = client.post(
        "/models/discover",
        json={
            "workbooks": [
                {"path": "/Users/k/Models/DEMOFIN_FY27.xlsx"},
                {"path": "/Users/k/Models/notes.xlsx", "text": ["ISIN INE000A01022"]},
                {"path": "/Users/k/Models/budget.xlsx"},
            ]
        },
    ).json()["candidates"]
    assert [(c["symbol"], c["method"]) for c in found] == [
        ("DEMOFIN", "file_name"),
        ("ACMEIND", "contents"),
        (None, "none"),
    ]
    linked = client.post(
        "/models/link",
        json={"symbol": "DEMOFIN", "workbook": {"path": "/Users/k/Models/DEMOFIN_FY27.xlsx"}},
    ).json()
    assert linked["origin"] == "discovered" and linked["doc_tag"]
    again = client.post(
        "/models/discover",
        json={
            "workbooks": [
                {
                    "path": "/elsewhere/renamed.xlsx",
                    "tag_model_id": linked["doc_tag"],
                    "tag_symbol": "DEMOFIN",
                }
            ]
        },
    ).json()["candidates"]
    assert again[0]["method"] == "tag" and again[0]["model_id"] == linked["id"]
    assert db.scalar(select(Filing).where(Filing.company_id == demo.id)) is not None


def test_storage_down_is_a_fast_503(client):
    from analekh.ingest.storage import StorageUnavailable

    def down():
        raise StorageUnavailable("object storage is not reachable at localhost:9000")

    app.dependency_overrides[get_store] = down
    r = client.get("/templates")
    assert r.status_code == 503
    assert "docker compose up -d minio" in r.json()["detail"]


def test_comps_export_shows_each_company_basis(client, db, store):
    solo = Company(nse_symbol="SOLO", name="Solo Systems Limited", isin="INE000S01011")
    db.add(solo)
    db.flush()
    q = pnl(500_000_000, 10_000_000, 400_000_000, 25_000_000)
    ingest_results_xbrl(db, store, solo, build_instance(**Q2, values_q=q, nature="Standalone"))
    ws = load_workbook(io.BytesIO(client.post("/comps", json={"symbols": ["SOLO"]}).content))
    assert ws["Comps"]["B5"].value == "Standalone"
    forced = client.post("/comps", json={"symbols": ["SOLO"], "consolidated": True}).content
    assert load_workbook(io.BytesIO(forced))["Comps"]["B5"].value == "Consolidated"


def test_glance_lines_changes_and_derived_ebitda(client, db, store):
    glow = Company(nse_symbol="GLOW", name="Glow Industries Limited", isin="INE000G01011")
    db.add(glow)
    db.flush()
    crore = 10_000_000

    def quarter(start, end, revenue, pat):
        q = {
            "RevenueFromOperations": revenue * crore,
            "OtherIncome": 5 * crore,
            "Income": (revenue + 5) * crore,
            "Expenses": (revenue - 20) * crore,
            "FinanceCosts": 2 * crore,
            "DepreciationDepletionAndAmortisationExpense": 3 * crore,
            "ProfitBeforeTax": 25 * crore,
            "TaxExpense": (25 - pat) * crore,
            "ProfitLossForPeriod": pat * crore,
        }
        return build_instance(q_start=start, q_end=end, ytd_start=start, values_q=q)

    ingest_results_xbrl(db, store, glow, quarter(date(2025, 7, 1), date(2025, 9, 30), 80, 16))
    ingest_results_xbrl(db, store, glow, quarter(date(2026, 4, 1), date(2026, 6, 30), 90, 18))
    ingest_results_xbrl(db, store, glow, quarter(date(2026, 7, 1), date(2026, 9, 30), 100, 20))
    g = client.get("/companies/GLOW/glance").json()
    assert g["quarter_label"] == "Q2 FY27" and g["previous_end"] == "2026-06-30"
    lines = {line["key"]: line for line in g["lines"]}
    assert float(lines["revenue"]["qoq"]) == 11.1 and float(lines["revenue"]["yoy"]) == 25.0
    # EBITDA = PBT 25 + finance 2 + depreciation 3 - other income 5 = 25 crore
    assert (
        float(lines["ebitda"]["value"]) == 25 * crore and "other income" in lines["ebitda"]["note"]
    )
    assert (
        float(lines["ebitda_margin"]["value"]) == 25.0 and lines["ebitda_margin"]["kind"] == "pct"
    )
    # Margins change in basis points, worked out from unrounded margins: EBITDA margin 25/100 =
    # 25.0% against 25/90 = 27.78% is -278 bps; net margin is 20% in all three quarters.
    assert lines["ebitda_margin"]["qoq"] == "-278" and lines["net_margin"]["qoq"] == "0"
    assert "exceptional_items" not in lines  # not tagged: no row

    # A lender's finance costs are its cost of business: no EBITDA lines for financial sectors.
    glow.sector = "Banking Services"
    db.flush()
    lender = {line["key"] for line in client.get("/companies/GLOW/glance").json()["lines"]}
    assert "ebitda" not in lender and "ebitda_margin" not in lender and "revenue" in lender


def test_glance_for_a_life_insurer_shows_premium_and_ratios(client, db, store):
    life = Company(nse_symbol="LIFEX", name="Lifex Life Insurance Limited", isin="INE000L01011")
    db.add(life)
    db.flush()
    crore = 10_000_000

    def quarter(start, end, net_premium, pat, solvency):
        q = {
            "GrossPremiumIncome": (net_premium + 30) * crore,
            "NetPremiumIncome": net_premium * crore,
            "ProfitLossBeforeTax": (pat + 5) * crore,
            "ProvisionsForTaxes": 5 * crore,
            "ProfitLossAfterTaxAndExtraordinaryItems": pat * crore,
            "SolvencyRatio": solvency,  # tagged as a fraction
            "PersistencyRatio13ThMonth": "0.8077",
            "ExpensesOfManagementRatio": "0.2038",
        }
        return build_instance(
            q_start=start, q_end=end, ytd_start=start, values_q=q, nature="Standalone"
        )

    ingest_results_xbrl(
        db, store, life, quarter(date(2026, 4, 1), date(2026, 6, 30), 2000, 28, "1.9803")
    )
    ingest_results_xbrl(
        db, store, life, quarter(date(2026, 7, 1), date(2026, 9, 30), 2500, 43, "1.8041")
    )
    assert life.results_taxonomy == "life_insurance"  # detected from the premium line items
    g = client.get("/companies/LIFEX/glance").json()
    lines = {line["key"]: line for line in g["lines"]}
    assert list(lines)[:2] == ["gross_premium", "revenue"]
    assert lines["revenue"]["label"] == "Net premium income"
    assert float(lines["revenue"]["qoq"]) == 25.0
    # Ratios read as percentages, and move in bps: 180.41% against 198.03% is -1,762 bps.
    assert (
        float(lines["solvency_ratio"]["value"]) == 180.4
        and lines["solvency_ratio"]["qoq"] == "-1762"
    )
    assert "150%" in lines["solvency_ratio"]["note"]
    assert "ebitda" not in lines and "persistency_61m" not in lines  # not tagged: no row


def test_year_ago_filed_as_a_half_year_explains_the_blank_yoy(client, db, store):
    # Canara Robeco AMC: its first results after listing (Sept 2025) cover April-September only.
    amc = Company(
        nse_symbol="HALFAMC",
        name="Half Year AMC Limited",
        isin="INE000H01011",
        sector="Half-Year AMC Test Sector",  # its own sector: real reporters fill page one
    )
    db.add(amc)
    db.flush()
    crore = 10_000_000
    q = pnl(120 * crore, 0, 60 * crore, 15 * crore)
    for start, end in (
        (date(2025, 4, 1), date(2025, 9, 30)),
        (date(2026, 7, 1), date(2026, 9, 30)),
    ):
        data = build_instance(
            q_start=start, q_end=end, ytd_start=start, values_q=q, nature="Standalone"
        )
        ingest_results_xbrl(db, store, amc, data)

    g = client.get("/companies/HALFAMC/glance").json()
    revenue = next(line for line in g["lines"] if line["key"] == "revenue")
    assert revenue["yoy"] is None
    assert g["year_ago_note"] == (
        "Q2 FY26 wasn't filed as a quarter: the standalone results for 30 Sep 2025 cover "
        "1 Apr 2025 to 30 Sep 2025 only"
    )
    assert g["previous_note"] == "No standalone results for Q1 FY27 on record"

    params = {"quarter_end": "2026-09-30", "sector": "Half-Year AMC Test Sector"}
    d = client.get("/desk", params=params).json()
    row = next(r for r in d["rows"] if r["symbol"] == "HALFAMC")
    assert row["revenue_yoy"] is None and row["yoy_note"] == g["year_ago_note"]


def test_missing_comparison_quarter_says_why(client, db, store):
    # Listed mid-year, first results for the half-year only (Canara Robeco AMC, September 2025):
    # the year-ago quarter was never filed on its own, so YoY is a dash with a reason.
    halfy = Company(
        nse_symbol="HALFY",
        name="Halfy Asset Management Limited",
        isin="INE000H01011",
        sector="Half-Year Test Sector",
    )
    db.add(halfy)
    db.flush()
    for start, end, revenue in (
        (date(2025, 4, 1), date(2025, 9, 30), 2_287_214_000),  # half-year
        (date(2026, 4, 1), date(2026, 6, 30), 1_162_014_000),
        (date(2026, 7, 1), date(2026, 9, 30), 1_215_889_000),
    ):
        q = pnl(revenue, 0, revenue // 2, revenue // 10)
        data = build_instance(q_start=start, q_end=end, ytd_start=start, values_q=q)
        ingest_results_xbrl(db, store, halfy, data)

    g = client.get("/companies/HALFY/glance").json()
    revenue = next(line for line in g["lines"] if line["key"] == "revenue")
    assert revenue["yoy"] is None and revenue["qoq"] is not None
    assert g["previous_note"] is None
    assert g["year_ago_note"].startswith("Q2 FY26 wasn't filed as a quarter")
    assert "1 Apr 2025 to 30 Sep 2025" in g["year_ago_note"]

    params = {"quarter_end": "2026-09-30", "sector": "Half-Year Test Sector"}
    row = client.get("/desk", params=params).json()["rows"]
    mine = next(r for r in row if r["symbol"] == "HALFY")
    assert mine["revenue_yoy"] is None and mine["yoy_note"] == g["year_ago_note"]


def test_results_desk_query_count_does_not_grow_with_companies(db, store):
    """The desk's database round trips are a fixed handful, however many companies reported (each
    costs milliseconds on the way to the database); the blank-year-ago notes come in one query."""
    from sqlalchemy import event

    from analekh.desk import desk

    sector = "Query Count Sector"

    def add(n: int) -> None:
        for i in range(n):
            c = Company(
                nse_symbol=f"QC{n}X{i}",
                name=f"Qc {n} {i} Ltd",
                isin=f"INE{n:03d}Q{i:02d}011",
                sector=sector,
            )
            db.add(c)
            db.flush()
            revenue = 1_000_000_000 + (n * 100 + i) * 1000  # each file unique
            q = pnl(revenue, 0, 500_000_000, 100_000_000)
            start, end = date(2026, 7, 1), date(2026, 9, 30)
            ingest_results_xbrl(
                db, store, c, build_instance(q_start=start, q_end=end, ytd_start=start, values_q=q)
            )

    def count() -> tuple[int, list[str | None]]:
        seen: list[str] = []
        conn = db.connection()

        def log(*args: object) -> None:
            seen.append("q")

        event.listen(conn, "before_cursor_execute", log)
        try:
            d = desk(db, quarter_end=date(2026, 9, 30), sector=sector, limit=50)
        finally:
            event.remove(conn, "before_cursor_execute", log)
        return len(seen), [r.yoy_note for r in d.rows]

    add(2)
    few, notes = count()
    assert (
        notes == ["No consolidated results for Q2 FY26 on record"] * 2
        or notes == ["No standalone results for Q2 FY26 on record"] * 2
    )
    add(6)
    many, notes = count()
    assert len(notes) == 8 and len(set(notes)) == 1
    assert many == few  # no query per company


def test_results_desk_rows_movers_and_sector_progress(client, db, store):
    from datetime import UTC, datetime

    crore = 10_000_000
    sector = "Desk Test Sector"

    def file_quarter(company, start, end, revenue, pat, filed):
        tax = 5 if pat > 0 else 0
        q = pnl(revenue * crore, 0, (revenue - pat - tax) * crore, tax * crore)
        data = build_instance(q_start=start, q_end=end, ytd_start=start, values_q=q)
        ingest_results_xbrl(db, store, company, data, filed_at=filed)

    companies = {}
    for symbol, isin in (
        ("GROWER", "INE000T01011"),
        ("SLIDER", "INE000T01029"),
        ("TINY", "INE000T01037"),
    ):
        c = Company(nse_symbol=symbol, name=f"{symbol.title()} Ltd", isin=isin, sector=sector)
        db.add(c)
        companies[symbol] = c
    db.flush()
    last_year = (date(2025, 7, 1), date(2025, 9, 30))
    this_year = (date(2026, 7, 1), date(2026, 9, 30))
    for symbol, (rev0, pat0, rev1, pat1, day) in {
        "GROWER": (100, 10, 120, 12, 12),  # +20% revenue, +20% profit
        "SLIDER": (200, 20, 180, -5, 13),  # -10% revenue, turned to a loss
        "TINY": (20, 2, 40, 4, 11),  # under Rs 50 crore this quarter: in the table, not movers
    }.items():
        c = companies[symbol]
        file_quarter(c, *last_year, rev0, pat0, datetime(2025, 10, day, 12, tzinfo=UTC))
        file_quarter(c, *this_year, rev1, pat1, datetime(2026, 10, day, 12, tzinfo=UTC))

    d = client.get("/desk", params={"quarter_end": "2026-09-30", "sector": sector}).json()
    assert d["quarter_label"] == "Q2 FY27" and d["reported"] == 3
    assert [r["symbol"] for r in d["rows"]] == ["SLIDER", "GROWER", "TINY"]  # newest filing first
    grower = next(r for r in d["rows"] if r["symbol"] == "GROWER")
    assert float(grower["revenue_yoy"]) == 20.0 and float(grower["pat_yoy"]) == 20.0
    # Only two companies qualify, so both are "top" and the bottom list stays empty.
    assert [r["symbol"] for r in d["revenue_top"]] == ["GROWER", "SLIDER"]
    assert d["revenue_bottom"] == []
    assert [r["symbol"] for r in d["turned_loss"]] == ["SLIDER"]
    assert all(r["symbol"] != "TINY" for r in d["revenue_top"] + d["revenue_bottom"])
    (progress,) = d["sectors"]
    assert progress["reported"] == 3 and progress["total"] == 3


def test_latest_close_on_the_company_page_and_largest_first_on_the_desk(client, db, store):
    crore = 10_000_000
    sector = "Price Test Sector"  # its own sector: real reporters fill the desk's pages
    big = Company(nse_symbol="BIGCAP", name="Big Cap Ltd", isin="INE000B01011", sector=sector)
    small = Company(nse_symbol="SMALLCAP", name="Small Cap Ltd", isin="INE000C01011", sector=sector)
    db.add_all([big, small])
    db.flush()
    for c, revenue in ((big, 900), (small, 100)):  # different files: one hash, one filing
        data = build_instance(
            q_start=date(2026, 7, 1),
            q_end=date(2026, 9, 30),
            ytd_start=date(2026, 4, 1),
            values_q=pnl(revenue * crore, 0, 60 * crore, 15 * crore),
        )
        ingest_results_xbrl(db, store, c, data)
    # Dates no real bhavcopy is loaded for (tests share the local database).
    db.add_all(
        [
            Price(company_id=big.id, trade_date=date(2020, 1, 2), close=900, market_cap=1),
            Price(
                company_id=big.id,
                trade_date=date(2020, 1, 3),
                close=1010,
                prev_close=1000,
                volume=5000,
                market_cap=90_000 * crore,
            ),
            Price(company_id=small.id, trade_date=date(2020, 1, 3), close=50, market_cap=crore),
        ]
    )
    db.flush()

    p = client.get("/companies/BIGCAP/price").json()
    assert p["trade_date"] == "2020-01-03"  # the newest stored day
    assert float(p["close"]) == 1010 and float(p["change_pct"]) == 1.0
    assert float(p["market_cap"]) == 90_000 * crore
    assert client.get("/companies/SMALLCAP/price").json()["change_pct"] is None  # no prev close

    params = {"quarter_end": "2026-09-30", "sector": sector, "sort": "market_cap"}
    rows = client.get("/desk", params=params).json()["rows"]
    assert [r["symbol"] for r in rows] == ["BIGCAP", "SMALLCAP"]
    assert rows[0]["price_date"] == "2020-01-03"
    assert client.get("/desk", params=params | {"sort": "size"}).status_code == 422


def test_peer_lists_and_the_comps_table(client, demo, db):
    db.execute(delete(PeerList))  # the owner's own lists, in a rolled-back transaction
    made = client.post("/peer-lists", json={"name": "  Gold  lenders ", "symbols": ["demofin"] * 2})
    assert made.status_code == 201
    assert made.json()["name"] == "Gold lenders" and made.json()["symbols"] == ["DEMOFIN"]
    assert client.post("/peer-lists", json={"name": "GOLD LENDERS"}).status_code == 409
    unknown = client.post("/peer-lists", json={"name": "Other", "symbols": ["NOPE"]})
    assert unknown.status_code == 422 and "NOPE" in unknown.json()["detail"]
    assert client.post("/peer-lists", json={"name": "   "}).status_code == 422
    list_id = made.json()["id"]
    updated = client.put(f"/peer-lists/{list_id}", json={"symbols": ["DEMOFIN", "ACMEIND"]})
    assert updated.json()["symbols"] == ["DEMOFIN", "ACMEIND"]
    assert [p["name"] for p in client.get("/peer-lists").json()] == ["Gold lenders"]

    # A symbol that has left the universe is reported, not fatal.
    table = client.post("/comps/table", json={"symbols": ["DEMOFIN", "GONECO"]}).json()
    assert table["missing"] == ["GONECO"]
    [row] = table["rows"]
    assert row["symbol"] == "DEMOFIN" and row["ttm_revenue"] is not None
    assert row["quarter_label"] and row["ev"] is None  # no EV until debt and cash are sourced
    assert set(table["stats"]) >= {"pe", "pb", "ebitda_margin", "revenue_yoy"}
    assert client.post("/comps/table", json={"symbols": [" "]}).status_code == 422
    sheet = client.post("/comps", json={"peer_list": "Gold lenders"})
    assert sheet.status_code == 200 and "comps.xlsx" in sheet.headers["content-disposition"]
    assert load_workbook(io.BytesIO(sheet.content))["Comps"]["A5"].value == "Demo Finance Limited"

    assert client.delete(f"/peer-lists/{list_id}").status_code == 204
    assert client.delete(f"/peer-lists/{list_id}").status_code == 404


def test_peer_suggestions_stay_in_sector(client, demo):
    peers = client.get("/companies/DEMOFIN/peers").json()
    assert all(p["sector"] == "Financial Services" and p["symbol"] != "DEMOFIN" for p in peers)
    assert client.get("/companies/ACMEIND/peers").json() == []  # no sector, no suggestions


def test_status_reads_without_contacting_nse(client, demo):
    s = client.get("/status").json()
    assert s["data"]["companies"] > 0 and s["watcher"]["next_ask_at"]
    assert s["prices"]["expected_day"] and isinstance(s["ai_configured"], bool)


def test_alert_channels_and_test_messages_never_leak_secrets(client, monkeypatch):
    from analekh.api.routers import system
    from analekh.config import get_settings

    s = get_settings()
    monkeypatch.setattr(s, "telegram_bot_token", "")
    monkeypatch.setattr(s, "smtp_url", "smtp://bot@mail.example.com:587")
    monkeypatch.setattr(s, "alert_email_to", "kaushal@example.com")
    channels = client.get("/alerts/channels").json()
    assert channels == {"telegram": False, "email": True, "email_to": "k•••@example.com"}
    off = client.post("/alerts/test", json={"channel": "telegram"}).json()
    assert not off["ok"] and "TELEGRAM_BOT_TOKEN" in off["detail"]

    monkeypatch.setattr(s, "telegram_bot_token", "SECRET-TOKEN")
    monkeypatch.setattr(s, "telegram_chat_id", "42")

    def refused(_text):
        raise RuntimeError("https://api.telegram.org/botSECRET-TOKEN/sendMessage: 401")

    monkeypatch.setattr(system, "send_telegram", refused)
    failed = client.post("/alerts/test", json={"channel": "telegram"}).json()
    assert not failed["ok"] and "SECRET" not in failed["detail"]
    monkeypatch.setattr(system, "send_email", lambda subject, body: True)
    sent = client.post("/alerts/test", json={"channel": "email"}).json()
    assert sent == {"ok": True, "detail": "Sent to k•••@example.com"}
    assert client.post("/alerts/test", json={"channel": "sms"}).status_code == 422


def test_comps_refuses_more_than_thirty_companies(client, demo):
    many = [f"CO{i}" for i in range(31)]
    too_many = client.post("/comps/table", json={"symbols": many})
    assert too_many.status_code == 422 and "at most 30" in too_many.json()["detail"]


def test_prices_are_overdue_only_after_the_evening_window():
    from datetime import UTC, datetime

    from analekh.api.routers.system import PRICES_DUE_AFTER
    from analekh.jobs.tasks import last_close_day

    def expected(h, m):  # IST on Friday 16 October 2026
        return last_close_day(datetime(2026, 10, 16, h, m, tzinfo=UTC) - PRICES_DUE_AFTER)

    # 23:00 IST = 17:30 UTC: the job may still run, so Thursday's close is what's expected.
    assert expected(17, 30).isoformat() == "2026-10-15"
    assert expected(18, 0).isoformat() == "2026-10-16"  # 23:30 IST: Friday's close is due


def test_results_history_and_its_excel_export(client, demo):
    h = client.get("/companies/DEMOFIN/history", params={"quarters": 4}).json()
    assert len(h["quarter_ends"]) == 4 and h["quarter_ends"] == sorted(h["quarter_ends"])
    assert h["quarter_labels"][-1].startswith("Q")
    revenue = next(line for line in h["lines"] if line["key"] == "revenue")
    assert len(revenue["values"]) == 4 and all(v is not None for v in revenue["values"])
    assert client.get("/companies/DEMOFIN/history", params={"quarters": 99}).status_code == 200
    sheet = client.get("/companies/DEMOFIN/history.xlsx")
    assert "DEMOFIN-results" in sheet.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(sheet.content))["Results"]
    assert ws["A4"].value == "Line" and ws["B4"].value.startswith("Q")


def test_watchlist_rows_carry_growth_and_price(client, demo, db):
    db.execute(delete(Watchlist))  # the owner's own watchlist, in a rolled-back transaction
    client.post("/watchlist", json={"symbol": "DEMOFIN"})
    [entry] = client.get("/watchlist").json()
    assert entry["quarter_label"] and entry["revenue"] is not None
    assert entry["revenue_yoy"] is not None  # eight quarters on record: a year-ago quarter exists
    assert entry["consolidated"] is not None


def test_analyse_fails_fast_when_the_job_queue_is_down(client, demo, monkeypatch):
    from analekh.jobs import health

    monkeypatch.setattr(health, "queue_available", lambda timeout=0.4, **_: False)
    r = client.post("/companies/DEMOFIN/analyse")
    assert r.status_code == 503 and "Background jobs aren't running" in r.json()["detail"]
    assert client.get("/system/jobs").json()["available"] is False


def test_calendar_endpoint_scopes_to_the_watchlist(client, db):
    from datetime import UTC, datetime

    from analekh.config import get_settings
    from analekh.db.models import FeedItem, Watchlist

    user = get_settings().default_user_id
    a = Company(nse_symbol="CEPAAA", name="Cep A Ltd", isin="INE000Y01011", sector="S")
    b = Company(nse_symbol="CEPBBB", name="Cep B Ltd", isin="INE000Y01029", sector="S")
    db.add_all([a, b])
    db.flush()
    db.add(Watchlist(user_id=user, company_id=b.id))
    when = datetime.now(UTC).date().replace(day=1)  # any day: the window below covers it
    for c, key in ((a, "ea"), (b, "eb")):
        db.add(
            FeedItem(
                feed="board_meetings",
                item_key=key,
                title=c.name,
                description=f"Financial Results |Meeting Date: {when.strftime('%d-%b-%Y')}",
                pub_date=datetime.now(UTC),
                company_id=c.id,
            )
        )
    db.flush()
    params = {"start": when.isoformat(), "end": when.isoformat()}
    everyone = client.get("/calendar", params={**params, "watchlist": "false"}).json()
    symbols = {
        e["symbol"]: e["watched"] for e in everyone["events"] if e["symbol"].startswith("CEP")
    }
    assert symbols == {"CEPAAA": False, "CEPBBB": True}
    mine = client.get("/calendar", params={**params, "watchlist": "true"}).json()
    assert {e["symbol"] for e in mine["events"]} >= {"CEPBBB"}
    assert "CEPAAA" not in {e["symbol"] for e in mine["events"]}
    assert (
        client.get("/calendar", params={"start": "2026-10-19", "end": "2024-01-01"}).status_code
        == 422
    )
