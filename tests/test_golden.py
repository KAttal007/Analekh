"""Golden-set worksheet: build, hand-check, import, run (analekh.golden; design doc 24)."""

from datetime import date

from openpyxl import load_workbook

from analekh import golden
from analekh.golden import Pick, fy_label, guess_unit
from tests.xbrl_factory import build_instance, pnl


def test_period_labels_follow_the_indian_financial_year():
    assert fy_label(date(2026, 4, 1), date(2026, 6, 30)) == "Q1 FY27"
    assert fy_label(date(2026, 1, 1), date(2026, 3, 31)) == "Q4 FY26"
    assert fy_label(date(2026, 4, 1), date(2026, 9, 30)) == "H1 FY27"
    assert fy_label(date(2025, 4, 1), date(2025, 12, 31)) == "9M FY26"
    assert fy_label(date(2025, 4, 1), date(2026, 3, 31)) == "FY26"


def test_unit_guess_from_xbrl_precision():
    assert guess_unit(-7) == guess_unit(-5) == "₹ crore"
    assert guess_unit(-4) == "₹ million"
    assert guess_unit(-3) == "₹ lakh"
    assert guess_unit(None) == "₹"


def q2_file() -> bytes:
    q = pnl(1_000_000_000, 50_000_000, 800_000_000, 62_500_000)
    q1 = pnl(900_000_000, 40_000_000, 720_000_000, 55_000_000)
    return build_instance(
        q_start=date(2026, 7, 1),
        q_end=date(2026, 9, 30),
        ytd_start=date(2026, 4, 1),
        values_q=q,
        values_ytd={k: q[k] + q1[k] for k in q},
    )


def test_marked_worksheet_becomes_a_manifest_that_golden_test_runs(tmp_path):
    xml = q2_file()
    sheet = golden.sheet_filing(
        xml, Pick("DEMO", date(2026, 9, 30), True, "Ind-AS"), "Demo Ltd", None
    )
    assert sheet.file == "demo_q2fy27_c.xml" and sheet.unit == "₹ crore"
    periods = {(x.line, x.period) for x in sheet.lines}
    assert {("revenue", "Q2 FY27"), ("revenue", "H1 FY27"), ("pat", "H1 FY27")} <= periods
    # Untagged quarter lines stay on the sheet for the analyst; nil exceptional items do not.
    gaps = {x.line for x in sheet.lines if x.fact is None}
    assert gaps == {"finance_costs", "depreciation", "eps_basic"}

    (tmp_path / sheet.file).write_bytes(xml)
    golden.write_workbook(tmp_path / golden.WORKSHEET, [sheet])
    wb = load_workbook(tmp_path / golden.WORKSHEET)
    values = wb["Values"]
    for row in values.iter_rows(min_row=2):
        line, period, found = row[12].value, row[5].value, row[16].value is not None
        if line == "other_income" and period == "Q2 FY27":
            row[8].value, row[9].value = "No", 5.1  # the PDF shows 5.10 crore, not 5.00
        elif not found:
            row[8].value = "Not in PDF" if line == "eps_basic" else None
        else:
            row[8].value = "Yes"
    wb.save(tmp_path / golden.WORKSHEET)

    summary = golden.import_sheet(tmp_path)
    assert (summary.corrected, summary.not_in_pdf, summary.pending) == (1, 1, 2)
    assert summary.confirmed == len(sheet.lines) - 4 and not summary.problems

    report = golden.run(tmp_path)
    assert report.checked == summary.confirmed + 1
    assert [(m.element, m.expected, m.actual) for m in report.mismatches] == [
        ("other_income", "51000000", "50000000")
    ]


def test_a_no_that_agrees_with_analekh_is_queried(tmp_path):
    xml = q2_file()
    sheet = golden.sheet_filing(xml, Pick("DEMO", date(2026, 9, 30), False, "Ind-AS"), "D", None)
    (tmp_path / sheet.file).write_bytes(xml)
    golden.write_workbook(tmp_path / golden.WORKSHEET, [sheet])
    wb = load_workbook(tmp_path / golden.WORKSHEET)
    row = next(r for r in wb["Values"].iter_rows(min_row=2) if r[12].value == "revenue")
    row[8].value, row[9].value = "No", 100.0  # 100 crore is exactly Analekh's value
    wb.save(tmp_path / golden.WORKSHEET)
    assert "agrees to two places" in golden.import_sheet(tmp_path).problems[0]
