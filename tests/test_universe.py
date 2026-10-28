from datetime import date
from decimal import Decimal

import httpx
import pytest
from sqlalchemy import delete, select

from analekh import coverage
from analekh.db.models import Company, Watchlist
from analekh.ingest.http import PoliteClient
from analekh.universe import sectors
from analekh.universe.loader import parse_equity_list, upsert_universe
from analekh.universe.search import search
from analekh.universe.sectors import apply_sectors, parse_constituents, refresh_sectors

# Column layout of NSE's EQUITY_L.csv (design doc 3.1). Companies are SYNTHETIC.
EQUITY_L = (
    b"SYMBOL,NAME OF COMPANY, SERIES, DATE OF LISTING, PAID UP VALUE, MARKET LOT, ISIN NUMBER,"
    b" FACE VALUE\n"
    b"""DEMOFIN,Demo Finance Limited,EQ,05-Jan-2012,10,1,INE000D01011,10
ACMEIND,Acme Industries Limited,EQ,17-Mar-2005,2,1,INE000A01022,2
OLDCO,Old Company Limited,BE,01-Jan-2000,10,1,INE000O01033,10
"""
)
SME_L = (
    b"SYMBOL,NAME_OF_COMPANY,SERIES,DATE_OF_LISTING,PAID_UP_VALUE,ISIN_NUMBER,FACE_VALUE,\n"
    b"""TINYSME,Tiny SME Limited,SM,12-Aug-24,10,INE000T01044,10,
"""
)


def test_parse_equity_list_handles_both_header_styles():
    rows = parse_equity_list(EQUITY_L) + parse_equity_list(SME_L, is_sme=True)
    demo = rows[0]
    assert demo.symbol == "DEMOFIN" and demo.series == "EQ" and demo.isin == "INE000D01011"
    assert demo.listing_date == date(2012, 1, 5) and demo.face_value == Decimal("10")
    sme = rows[-1]
    assert sme.symbol == "TINYSME" and sme.is_sme and sme.listing_date == date(2024, 8, 12)


def test_sector_worksheet_applies_only_confirmed_nse_sectors(db, tmp_path):
    from openpyxl import Workbook

    upsert_universe(db, parse_equity_list(EQUITY_L))
    wb = Workbook()
    ws = wb.active
    ws.append(["Symbol", "Company", "Sector", "Confidence", "Confirmed"])
    ws.append(["DEMOFIN", "Demo Finance", "Financial Services", "High", "Yes"])
    ws.append(["ACMEIND", "Acme Industries", "Capital Goods", "Low", None])  # not reviewed
    ws.append(["OLDCO", "Old Company", "Widgets", "Low", "Yes"])  # not one of NSE's sectors
    wb.save(tmp_path / "sectors.xlsx")
    map_path = tmp_path / "sector_map.csv"

    r = sectors.import_sector_worksheet(db, tmp_path / "sectors.xlsx", map_path)
    assert r == {"rows": 3, "confirmed": 2, "not_an_nse_sector": 1, "updated": 1}
    demo, acme = (db.scalar(select(Company).where(Company.nse_symbol == s)) for s in ("DEMOFIN", "ACMEIND"))
    assert (demo.sector, demo.sector_source) == ("Financial Services", "user")
    assert acme.sector is None
    assert sectors.load_sector_map(map_path) == {"DEMOFIN": ("Financial Services", None)}


def test_rights_entitlements_are_not_companies():
    data = EQUITY_L + b"DEMOFIN-RE,Demo Finance Limited-RE,BE,12-Oct-2026,10,1,INE000D20011,10\n"
    assert "DEMOFIN-RE" not in {r.symbol for r in parse_equity_list(data)}


def test_upsert_by_isin_renames_and_deactivates(db):
    first = upsert_universe(db, parse_equity_list(EQUITY_L))
    assert sorted(first.inserted) == ["ACMEIND", "DEMOFIN", "OLDCO"]

    week2 = EQUITY_L.replace(b"DEMOFIN,", b"DEMOFINANCE,").replace(b"OLDCO,Old Company Limited,BE,01-Jan-2000,10,1,INE000O01033,10\n", b"")
    second = upsert_universe(db, parse_equity_list(week2))
    assert second.renamed == [("DEMOFIN", "DEMOFINANCE")]
    assert second.deactivated == ["OLDCO"]
    old = db.scalar(select(Company).where(Company.nse_symbol == "OLDCO"))
    assert old is not None and not old.is_active
    assert db.scalar(select(Company).where(Company.isin == "INE000D01011")).nse_symbol == "DEMOFINANCE"


def test_search_symbol_prefix_and_fuzzy_name(db):
    upsert_universe(db, parse_equity_list(EQUITY_L))
    assert search(db, "demofin")[0].nse_symbol == "DEMOFIN"
    assert search(db, "acme industr")[0].nse_symbol == "ACMEIND"
    assert search(db, "INE000A01022")[0].nse_symbol == "ACMEIND"
    assert search(db, "") == []


def test_search_matches_words_in_any_order(db):
    names = (
        EQUITY_L.split(b"\n", 2)[0]
        + b"""
ZEDMO,Zed Motors Limited,EQ,01-Jan-2020,2,1,INE000Z01011,2
ZEDPV,Zed Motors Passenger Vehicles Limited,EQ,01-Jan-2000,2,1,INE000Z01022,2
QUICKMOT,Quick Motors Ltd,EQ,01-Jan-2010,10,1,INE000Q01033,10
ZEDAMC,Zed Asset Management Company Limited,EQ,01-Jan-2015,5,1,INE000Z01044,5
ZEDSTEEL,Zed Steel Limited,EQ,01-Jan-2005,1,1,INE000Z01055,1
"""
    )
    upsert_universe(db, parse_equity_list(names))

    def symbols(q: str) -> list[str]:
        return [h.nse_symbol for h in search(db, q)]

    assert symbols("zed pass") == ["ZEDPV"]  # skipped word, partial last word
    assert symbols("passenger zed") == ["ZEDPV"]  # any order
    assert symbols("zed motors") == ["ZEDMO", "ZEDPV"]  # not Quick Motors
    assert symbols("zed amc")[0] == "ZEDAMC"  # words joined as a symbol prefix
    assert set(symbols("ZED")) == {"ZEDAMC", "ZEDMO", "ZEDPV", "ZEDSTEEL"}  # not Quick Motors
    assert symbols("zed stel")[0] == "ZEDSTEEL"  # typo falls back to trigram similarity
    assert symbols("  zed   pass ") == ["ZEDPV"]
    assert symbols("zed (pass") == []  # regex characters are escaped, not interpreted
    assert symbols("zed_") == [] and symbols("z%") == []  # LIKE wildcards too


def test_search_narrows_with_each_keystroke(db):
    names = (
        EQUITY_L.split(b"\n", 2)[0]
        + b"""
TCS,Tazza Consultancy Services Limited,EQ,01-Jan-2004,1,1,INE000T01011,1
TCIEXP,TCI Express Limited,EQ,01-Jan-2016,2,1,INE000T01022,2
TIGER,Tiger Foods Limited,EQ,01-Jan-2010,10,1,INE000T01033,10
MATCHBX,Matchbox Industries Limited,EQ,01-Jan-2012,10,1,INE000M01044,10
OTCX,Orion Trading Company Limited,EQ,01-Jan-2018,10,1,INE000O01055,10
ATLAS,Atlas Tyres Limited,EQ,01-Jan-2001,10,1,INE000A01066,10
"""
    )
    upsert_universe(db, parse_equity_list(names))

    def symbols(q: str, **kw: int) -> list[str]:
        return [h.nse_symbol for h in search(db, q, **kw)]

    assert symbols("T") == ["TCIEXP", "TCS", "TIGER"]
    assert symbols("TC") == ["TCIEXP", "TCS", "MATCHBX", "OTCX"]
    assert symbols("tcs") == ["TCS"]
    assert symbols("T", limit=2) + symbols("T", limit=2, offset=2) == ["TCIEXP", "TCS", "TIGER"]
    assert symbols("T", limit=2, offset=3) == []


def test_sector_map_and_taxonomy(db):
    upsert_universe(db, parse_equity_list(EQUITY_L))
    apply_sectors(db, {"DEMOFIN": ("Financial Services", "NBFC")})
    demo = db.scalar(select(Company).where(Company.nse_symbol == "DEMOFIN"))
    assert demo.sector == "Financial Services" and demo.results_taxonomy == "nbfc"


def test_watchlist_controls_coverage_level(db):
    db.execute(delete(Watchlist))  # the owner's own watchlist, in a rolled-back transaction
    upsert_universe(db, parse_equity_list(EQUITY_L))
    demo = db.scalar(select(Company).where(Company.nse_symbol == "DEMOFIN"))
    assert coverage.level(db, demo) == "universe"
    coverage.add(db, demo, settings={"alerts": True})
    coverage.add(db, demo, settings={"alerts": False})  # idempotent upsert
    assert coverage.level(db, demo) == "deep"
    assert [w.settings for w in coverage.entries(db)] == [{"alerts": False}]
    coverage.remove(db, demo)
    assert not coverage.is_watched(db, demo.id)


NIFTY_HEADER = b"Company Name,Industry,Symbol,Series,ISIN Code\n"


def test_parse_constituents_ignores_header_case_and_bom():
    data = ("﻿company name , INDUSTRY,symbol,Series,ISIN Code\nAcme Ltd,Capital Goods,acme,EQ,X\n").encode()
    assert parse_constituents(data) == {"ACME": ("Capital Goods", None)}
    old = NIFTY_HEADER + b"Old Label Ltd,PHARMA,OLDLBL,EQ,X\nNew Ltd,Healthcare,NEWCO,EQ,Y\n"
    assert parse_constituents(old) == {"NEWCO": ("Healthcare", None)}


def test_refresh_sectors_from_index_lists(db, monkeypatch: pytest.MonkeyPatch):
    universe = (
        EQUITY_L
        + b"""ZEDBANK,Zed Bank Limited,EQ,01-Jan-2001,10,1,INE000Z01077,10
ZEDLIFE,Zed Life Insurance Limited,EQ,01-Jan-2017,10,1,INE000Z01088,10
"""
    )
    upsert_universe(db, parse_equity_list(universe))
    oldco = db.scalar(select(Company).where(Company.nse_symbol == "OLDCO"))
    oldco.sector = "Set by analyst"
    monkeypatch.setattr(sectors, "load_sector_map", lambda: {})

    lists = {
        "ind_niftytotalmarket_list.csv": NIFTY_HEADER
        + b"""Demo Finance Limited,Financial Services,DEMOFIN,EQ,INE000D01011
Acme Industries Limited,Capital Goods,ACMEIND,EQ,INE000A01022
Old Company Limited,Capital Goods,OLDCO,EQ,INE000O01033
Zed Bank Limited,Financial Services,ZEDBANK,EQ,INE000Z01077
Zed Life Insurance Limited,Financial Services,ZEDLIFE,EQ,INE000Z01088
""",
        "ind_niftybanklist.csv": NIFTY_HEADER + b"Zed Bank Limited,Financial Services,ZEDBANK,EQ,INE000Z01077\n",
    }

    def handler(request: httpx.Request) -> httpx.Response:
        name = request.url.path.rsplit("/", 1)[-1]
        return httpx.Response(200, content=lists[name]) if name in lists else httpx.Response(404)

    client = PoliteClient(transport=httpx.MockTransport(handler), backoff_base=0)
    result = refresh_sectors(db, client)

    assert result["lists_fetched"] == 2 and result["lists_total"] == 13  # missing lists skipped
    assert result["companies_listed"] == 5 and result["banks"] == 1
    by_symbol = {c.nse_symbol: c for c in db.scalars(select(Company))}
    assert by_symbol["ACMEIND"].sector == "Capital Goods"
    assert by_symbol["OLDCO"].sector == "Set by analyst"  # never overwritten
    assert by_symbol["ZEDBANK"].results_taxonomy == "banking"  # bank index membership
    assert by_symbol["ACMEIND"].results_taxonomy is None
    assert by_symbol["ZEDLIFE"].results_taxonomy is None


def test_stocks_workbook_import_and_source_precedence(db, tmp_path):
    from openpyxl import Workbook

    upsert_universe(db, parse_equity_list(EQUITY_L))
    by_symbol = {c.nse_symbol: c for c in db.scalars(select(Company))}
    by_symbol["OLDCO"].sector, by_symbol["OLDCO"].sector_source = "Analyst's call", "user"

    wb = Workbook()
    ws = wb.active
    ws.append(["Sector", "BSE Stock ID", "Security Code", "ISIN No"])
    ws.append(["Banking Services", "#VALUE!", "500001", "INE000D01011"])  # DEMOFIN
    ws.append(["Machinery, Equipment & Components", "#VALUE!", "500002", "INE000A01022"])  # ACME
    ws.append(["Chemicals", "#VALUE!", "500003", "INE000O01033"])  # OLDCO: user value wins
    ws.append(["#FIELD!", "#VALUE!", "500004", "INE999X01099"])  # error: skipped
    ws.append(["Food & Tobacco", "#VALUE!", "500005", "INE999Y01099"])  # not in the universe
    path = tmp_path / "stocks.xlsx"
    wb.save(path)

    result = sectors.import_stocks_workbook(db, path)
    assert result == {"rows": 5, "with_sector": 4, "matched": 3, "updated": 2}
    assert (by_symbol["DEMOFIN"].sector, by_symbol["DEMOFIN"].sector_source) == (
        "Banking Services",
        "excel_stocks",
    )
    assert by_symbol["DEMOFIN"].results_taxonomy is None  # LSEG "Banking Services" includes NBFCs
    assert by_symbol["OLDCO"].sector == "Analyst's call"

    assert by_symbol["DEMOFIN"].sector_lseg == "Banking Services"  # always kept

    assert apply_sectors(db, {"ACMEIND": ("Capital Goods", None)}, source="nse_indices") == 1
    assert (by_symbol["ACMEIND"].sector, by_symbol["ACMEIND"].sector_source) == (
        "Capital Goods",
        "nse_indices",
    )
    assert by_symbol["ACMEIND"].sector_lseg == "Machinery, Equipment & Components"
    assert apply_sectors(db, {"ACMEIND": ("Industrial Machinery", None)}, source="excel_stocks") == 0
    assert by_symbol["ACMEIND"].sector == "Capital Goods"


def test_stocks_workbook_needs_sector_and_isin(db, tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.append(["Name", "Ticker"])
    wb.save(tmp_path / "bad.xlsx")
    with pytest.raises(ValueError, match="needs Sector and ISIN No"):
        sectors.import_stocks_workbook(db, tmp_path / "bad.xlsx")
