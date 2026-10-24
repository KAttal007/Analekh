"""Golden-set accuracy (design doc 24): hand-checked values from real filings, compared exactly.

Layout: ``tests/fixtures/golden/manifest.yaml`` lists filings and the values an analyst checked
against the published results PDF::

    filings:
      - file: muthootfin_q4fy26_s.xml        # XBRL instance, next to the manifest
        taxonomy: indas
        expect:
          - line: revenue                    # resolved through xbrl/elements.py, so the
            period_end: 2026-03-31           # element tables are tested too
            period_start: 2026-01-01         # omit for balance-sheet instants
            value: "81799230000"
          - element: OtherIncome             # or a raw element name
            ...

The manifest comes from a worksheet: ``analekh golden-sheet`` copies the picked filings' XBRL
from storage (no downloads) and writes ``golden-worksheet.xlsx`` with each line, Analekh's value
and a link to the results PDF; the analyst marks every line against the PDF, and
``analekh golden-import`` writes the manifest. A confirmed line expects the XBRL value (the PDF is
rounded); a corrected one expects the PDF's value, so it fails until the cause is fixed.

Raw exchange files are not committed if licensing is unclear (section 23): keep the folder
local, or in a private CI artifact. Target: 100% exact match on mapped elements.
"""

import json
from dataclasses import asdict, dataclass, field
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any
from urllib.parse import quote_plus

import yaml
from sqlalchemy import select
from sqlalchemy.orm import Session

from analekh.ingest.storage import ObjectStore
from analekh.xbrl.elements import candidates
from analekh.xbrl.parser import ParsedFact, ParsedFiling, parse_instance

DEFAULT_DIR = Path("tests/fixtures/golden")
WORKSHEET = "golden-worksheet.xlsx"


@dataclass
class Mismatch:
    file: str
    element: str
    period_end: str
    expected: str
    actual: str | None


@dataclass
class GoldenReport:
    fixtures: int = 0
    checked: int = 0
    matched: int = 0
    mismatches: list[Mismatch] = field(default_factory=list)
    missing_files: list[str] = field(default_factory=list)
    note: str = ""

    @property
    def accuracy(self) -> float | None:
        return self.matched / self.checked if self.checked else None

    @property
    def passed(self) -> bool:
        return self.checked > 0 and self.matched == self.checked and not self.missing_files

    def as_dict(self) -> dict[str, Any]:
        return asdict(self) | {"accuracy": self.accuracy, "passed": self.passed}


def resolve(
    filing: ParsedFiling, line: str, taxonomy: str, start: date | None, end: date
) -> ParsedFact | None:
    """A logical line's fact for a period, through the taxonomy's element table."""
    for name in candidates(line, taxonomy):
        fact = filing.get(name, end=end, start=start, instant=None if start else True)
        if fact is not None:
            return fact
    return None


def run(folder: Path = DEFAULT_DIR) -> GoldenReport:
    manifest = folder / "manifest.yaml"
    report = GoldenReport()
    if not manifest.exists():
        report.note = f"no golden set at {manifest}; add hand-checked filings (Phase 0/1)"
        return report
    spec = yaml.safe_load(manifest.read_text(encoding="utf-8")) or {}
    for entry in spec.get("filings", []):
        path = folder / entry["file"]
        if not path.exists():
            report.missing_files.append(entry["file"])
            continue
        report.fixtures += 1
        filing = parse_instance(path.read_bytes())
        taxonomy = entry.get("taxonomy") or filing.taxonomy()
        for exp in entry.get("expect", []):
            report.checked += 1
            end = date.fromisoformat(str(exp["period_end"]))
            start = (
                date.fromisoformat(str(exp["period_start"])) if exp.get("period_start") else None
            )
            if exp.get("line"):
                fact = resolve(filing, exp["line"], taxonomy, start, end)
            else:
                fact = filing.get(
                    exp["element"], end=end, start=start, instant=None if start else True
                )
            expected = Decimal(str(exp["value"]))
            if fact is not None and fact.value == expected:
                report.matched += 1
            else:
                report.mismatches.append(
                    Mismatch(
                        entry["file"],
                        exp.get("line") or exp["element"],
                        str(end),
                        str(expected),
                        str(fact.value) if fact else None,
                    )
                )
    return report


def write(report: GoldenReport, out: Path) -> None:
    out.write_text(json.dumps(report.as_dict(), indent=2, default=str))


# --- Worksheet: the hand-check (design doc 24) ---------------------------------------------------


@dataclass(frozen=True)
class Pick:
    symbol: str
    period_end: date
    consolidated: bool
    format: str  # as the worksheet names it


def _pick(symbol: str, end: str, statement: str, fmt: str) -> Pick:
    return Pick(symbol, date.fromisoformat(end), statement == "C", fmt)


# Twenty filings across every results format, established large caps only. Recent filings are
# picked from those NSE gave a results PDF link for; four from September 2023 cover the schemas
# used before integrated filings (Ind-AS 2020, in-bse-fin and the RBI entry point), whose PDFs
# are on the company's site. Quarter ends vary so the year to date is a half year, nine months or
# a full year.
PICKS: tuple[Pick, ...] = (
    _pick("BHARTIARTL", "2025-06-30", "C", "Ind-AS"),
    _pick("M&M", "2025-09-30", "C", "Ind-AS"),
    _pick("NTPC", "2025-03-31", "C", "Ind-AS"),
    _pick("LT", "2025-03-31", "C", "Ind-AS"),
    _pick("TITAN", "2023-09-30", "C", "Ind-AS"),
    _pick("ITC", "2026-06-30", "S", "Ind-AS"),
    _pick("ULTRACEMCO", "2025-09-30", "S", "Ind-AS"),
    _pick("HINDUNILVR", "2025-06-30", "S", "Ind-AS"),
    _pick("BEL", "2023-09-30", "S", "Ind-AS"),
    _pick("ICICIBANK", "2025-12-31", "C", "Bank"),
    _pick("KOTAKBANK", "2026-03-31", "S", "Bank"),
    _pick("SBIN", "2026-06-30", "S", "Bank"),
    _pick("AXISBANK", "2023-09-30", "S", "Bank"),
    _pick("BAJFINANCE", "2025-03-31", "C", "NBFC"),
    _pick("LICHSGFIN", "2025-06-30", "S", "NBFC"),
    _pick("CHOLAFIN", "2023-09-30", "S", "NBFC"),
    _pick("NIACL", "2026-03-31", "S", "General insurer"),
    _pick("STARHEALTH", "2025-12-31", "S", "General insurer"),
    _pick("LICI", "2025-09-30", "S", "Life insurer"),
    _pick("HDFCLIFE", "2026-06-30", "S", "Life insurer"),
)

# Lines checked for the quarter, by results format, as the PDF names them. Revenue and profit
# are also checked for the year to date (results XBRL tags only the current period and the year
# to date, not last year's columns): the dates, not the context ids, pick the column (rule 5).
LINES: dict[str, tuple[tuple[str, str], ...]] = {
    "indas": (
        ("revenue", "Revenue from operations"),
        ("other_income", "Other income"),
        ("total_income", "Total income"),
        ("total_expenses", "Total expenses"),
        ("finance_costs", "Finance costs"),
        ("depreciation", "Depreciation and amortisation"),
        ("exceptional_items", "Exceptional items"),
        ("pbt", "Profit before tax"),
        ("tax", "Tax expense"),
        ("pat", "Profit for the period"),
        ("eps_basic", "Basic EPS"),
    ),
    "banking": (
        ("revenue", "Interest earned"),
        ("other_income", "Other income"),
        ("total_income", "Total income"),
        ("interest_expended", "Interest expended"),
        ("operating_expenses", "Operating expenses"),
        ("operating_profit", "Operating profit before provisions"),
        ("provisions", "Provisions (other than tax) and contingencies"),
        ("pbt", "Profit before tax"),
        ("tax", "Tax expense"),
        ("pat", "Net profit for the period"),
        ("eps_basic", "Basic EPS"),
        ("gross_npa", "Gross NPA (amount)"),
        ("net_npa", "Net NPA (amount)"),
    ),
    "insurance": (
        ("gross_premium", "Gross premium written"),
        ("net_premium", "Net premium written"),
        ("revenue", "Premium earned (net)"),
        ("investment_income", "Income from investments (revenue account)"),
        ("claims_incurred", "Claims incurred (net)"),
        ("operating_profit", "Operating profit / (loss)"),
        ("pbt", "Profit before tax"),
        ("tax", "Provision for tax"),
        ("pat", "Profit after tax"),
        ("eps_basic", "Basic EPS"),
    ),
    "life_insurance": (
        ("gross_premium", "Gross premium income"),
        ("revenue", "Net premium income"),
        ("investment_income", "Income from investments (policyholders)"),
        ("total_income", "Total income (policyholders' account)"),
        ("total_expenses", "Total expenses (policyholders' account)"),
        ("surplus", "Surplus / (deficit)"),
        ("pbt", "Profit before tax"),
        ("tax", "Provision for tax"),
        ("pat", "Profit after tax"),
        ("eps_basic", "Basic EPS"),
        ("solvency_ratio", "Solvency ratio (as tagged)"),
    ),
}
COMPARED = ("revenue", "pat")

# What a PDF states its figures in; the worksheet guesses from the XBRL precision and the analyst
# corrects it from the PDF's header ("₹ in crore").
UNITS: dict[str, int] = {
    "₹": 1,
    "₹ thousand": 10**3,
    "₹ lakh": 10**5,
    "₹ million": 10**6,
    "₹ crore": 10**7,
}
ANSWERS = ("Yes", "No", "Not in PDF")


def guess_unit(decimals: int | None) -> str:
    """The PDF's unit from the XBRL precision: -5 is crore to two places (or lakh, whole), so
    this is a guess the analyst confirms."""
    if decimals is None or decimals >= 0:
        return "₹"
    if decimals <= -5:
        return "₹ crore"
    return {-4: "₹ million", -3: "₹ lakh"}.get(decimals, "₹ thousand")


def fy_label(start: date | None, end: date) -> str:
    """Q1 FY27 for April-June 2026; H1, 9M or FY for a year to date."""
    fy = end.year + 1 if end.month > 3 else end.year
    days = (end - start).days + 1 if start is not None else 0
    for least, label in ((300, "FY"), (240, "9M FY"), (150, "H1 FY")):
        if days > least:
            return f"{label}{fy % 100:02d}"
    quarter = {6: 1, 9: 2, 12: 3, 3: 4}.get(end.month, 0)
    return f"Q{quarter} FY{fy % 100:02d}"


@dataclass
class SheetLine:
    line: str
    label: str
    period: str
    start: date | None
    end: date
    fact: ParsedFact | None  # None: Analekh finds no value, so the analyst records the PDF's

    @property
    def kind(self) -> str:
        if self.fact is None:
            return {"eps_basic": "per_share", "solvency_ratio": "ratio"}.get(self.line, "money")
        unit = (self.fact.unit or "").lower()
        if "share" in unit:
            return "per_share"
        if unit == "pure":
            return "ratio"
        return "money"


def sheet_lines(filing: ParsedFiling, taxonomy: str) -> list[SheetLine]:
    """The lines to check in one filing: the quarter, and revenue and profit for the year to
    date."""
    start, end = filing.declared_start, filing.declared_end
    if start is None or end is None:
        return []
    table = LINES.get(taxonomy, LINES["indas"])
    compared = tuple(row for row in table if row[0] in COMPARED)
    periods: list[tuple[str, date, date, tuple[tuple[str, str], ...]]] = [
        (fy_label(start, end), start, end, table)
    ]
    fy_start = date(end.year if end.month > 3 else end.year - 1, 4, 1)
    if fy_start < start:
        periods.append((fy_label(fy_start, end), fy_start, end, compared))
    out: list[SheetLine] = []
    for period, p_start, p_end, rows in periods:
        for line, label in rows:
            fact = resolve(filing, line, taxonomy, p_start, p_end)
            # A quarter line Analekh cannot find stays on the sheet, so the gap is recorded; not
            # for exceptional items (often untagged when nil) or the year to date (older files
            # tag none).
            quarter_gap = p_start == start and line != "exceptional_items"
            if fact is not None or quarter_gap:
                out.append(SheetLine(line, label, period, p_start, p_end, fact))
    return out


@dataclass
class SheetFiling:
    file: str  # the XBRL copy next to the worksheet
    company: str
    symbol: str
    format: str
    consolidated: bool
    quarter: str
    audited: bool | None
    pdf_url: str | None
    unit: str
    taxonomy: str
    lines: list[SheetLine]
    filing_id: int | None = None


def file_name(symbol: str, quarter: str, consolidated: bool) -> str:
    slug = symbol.lower().replace("&", "and")
    return f"{slug}_{quarter.replace(' ', '').lower()}_{'c' if consolidated else 's'}.xml"


def sheet_filing(
    xml: bytes, pick: Pick, company: str, pdf_url: str | None, filing_id: int | None = None
) -> SheetFiling:
    parsed = parse_instance(xml)
    taxonomy = parsed.taxonomy()
    lines = sheet_lines(parsed, taxonomy)
    quarter = fy_label(parsed.declared_start, parsed.declared_end or pick.period_end)
    money = next((x.fact for x in lines if x.fact and x.kind == "money"), None)
    return SheetFiling(
        file=file_name(pick.symbol, quarter, pick.consolidated),
        company=company,
        symbol=pick.symbol,
        format=pick.format,
        consolidated=pick.consolidated,
        quarter=quarter,
        audited=parsed.audited,
        pdf_url=pdf_url,
        unit=guess_unit(money.decimals if money else None),
        taxonomy=taxonomy,
        lines=lines,
        filing_id=filing_id,
    )


def build_sheet(
    s: Session,
    store: ObjectStore,
    folder: Path = DEFAULT_DIR,
    picks: tuple[Pick, ...] = PICKS,
    force: bool = False,
) -> tuple[Path, list[SheetFiling], list[Pick]]:
    """Copy the picked filings' XBRL from storage and write the worksheet. Never downloads."""
    from analekh.db.models import Company, Filing

    out = folder / WORKSHEET
    if out.exists() and not force:
        raise FileExistsError(f"{out} exists and may hold your answers; pass --force to replace it")
    folder.mkdir(parents=True, exist_ok=True)
    built: list[SheetFiling] = []
    missing: list[Pick] = []
    for pick in picks:
        row = s.execute(
            select(Filing, Company.name)
            .join(Company, Company.id == Filing.company_id)
            .where(
                Company.nse_symbol == pick.symbol,
                Filing.kind == "results_xbrl",
                Filing.period_type == "Q",
                Filing.period_end == pick.period_end,
                Filing.consolidated.is_(pick.consolidated),
            )
            .order_by(Filing.filed_at.desc().nulls_last(), Filing.id.desc())
            .limit(1)
        ).first()
        if row is None:
            missing.append(pick)
            continue
        filing, company = row
        xml = store.get(filing.storage_key)
        pdf = ((filing.meta or {}).get("links") or {}).get("pdf")
        sheet = sheet_filing(xml, pick, company, pdf, filing.id)
        (folder / sheet.file).write_bytes(xml)
        built.append(sheet)
    write_workbook(out, built)
    return out, built, missing


# Values sheet columns; those after "Notes" are hidden and read back by the import.
VALUE_COLUMNS = (
    ("#", 5),
    ("Symbol", 13),
    ("Statement", 13),
    ("Filing", 10),
    ("Line", 44),
    ("Period", 22),
    ("Analekh value", 18),
    ("Unit", 13),
    ("Matches the PDF?", 17),
    ("PDF shows (if No)", 18),
    ("Notes", 40),
    ("File", 0),
    ("Line key", 0),
    ("Element", 0),
    ("Period start", 0),
    ("Period end", 0),
    ("Value (₹)", 0),
    ("Kind", 0),
    ("Taxonomy", 0),
)
FILING_COLUMNS = (
    ("File", 26),
    ("Company", 40),
    ("Symbol", 13),
    ("Format", 16),
    ("Statement", 13),
    ("Quarter", 10),
    ("Audited", 9),
    ("Results PDF", 12),
    ("PDF unit", 12),
    ("Marked", 9),
    ("Lines", 7),
    ("Filing id", 9),
)
HOW_TO = (
    "Golden set: hand-check of Analekh's values against the published results PDF",
    "",
    "1. Filings sheet: open a filing's results PDF. Read the unit in the table header (₹ in crore,"
    " lakh, million) and set PDF unit if the guess is wrong; Analekh's values convert.",
    "2. Statement: S is standalone, C consolidated. Integrated filings carry both tables; read"
    " the right one.",
    "3. Values sheet: for each line, compare Analekh's value with the PDF's figure for that"
    " period (the quarter, or the year to date: H1, 9M or FY).",
    "   Yes: they agree to the PDF's last digit (the PDF is rounded).",
    "   No: type the PDF's figure, in the PDF's unit, under 'PDF shows'.",
    "   Not in PDF: the PDF does not show this line.",
    "4. Notes: anything odd (a restated comparative, a line named differently).",
    "",
    "Basic EPS is not annualised for a quarter. NPAs are amounts, not percentages. Ratios are"
    " shown as tagged (1.85 can mean 185%).",
    "",
    "When done: save, close Excel, then run",
    "   uv run analekh golden-import",
    "   uv run analekh golden-test",
)


def write_workbook(path: Path, filings: list[SheetFiling]) -> None:
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    how = wb.active
    if how is None:
        raise RuntimeError("new workbook has no sheet")
    how.title = "How to check"
    for i, text in enumerate(HOW_TO, start=1):
        how.cell(i, 1, text).font = Font(bold=i == 1, size=13 if i == 1 else 11)
    how.column_dimensions["A"].width = 120

    lists = wb.create_sheet("Lists")
    for i, (name, factor) in enumerate(UNITS.items(), start=1):
        lists.cell(i, 1, name)
        lists.cell(i, 2, factor)
    for i, answer in enumerate(ANSWERS, start=1):
        lists.cell(i, 4, answer)
    lists.sheet_state = "hidden"
    unit_range = f"Lists!$A$1:$A${len(UNITS)}"
    factor_table = f"Lists!$A$1:$B${len(UNITS)}"

    header = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F3A5F")
    entry_fill = PatternFill("solid", fgColor="FFF8DC")

    fs = wb.create_sheet("Filings")
    vs = wb.create_sheet("Values")
    for sheet, columns in ((fs, FILING_COLUMNS), (vs, VALUE_COLUMNS)):
        for c, (title, width) in enumerate(columns, start=1):
            cell = sheet.cell(1, c, title)
            cell.font, cell.fill = header, header_fill
            cell.alignment = Alignment(vertical="center")
            letter = cell.column_letter
            if width:
                sheet.column_dimensions[letter].width = width
            else:
                sheet.column_dimensions[letter].hidden = True
        sheet.freeze_panes = "A2"

    unit_dv = DataValidation(type="list", formula1=unit_range, allow_blank=False)
    fs.add_data_validation(unit_dv)
    for r, f in enumerate(filings, start=2):
        statement = "C" if f.consolidated else "S"
        audited = "" if f.audited is None else ("Yes" if f.audited else "No")
        for c, value in enumerate(
            (f.file, f.company, f.symbol, f.format, statement, f.quarter, audited), start=1
        ):
            fs.cell(r, c, value)
        # No PDF link from NSE: a web search for the results, which the company's site carries.
        link = fs.cell(r, 8, "Open PDF" if f.pdf_url else "Search")
        link.hyperlink = f.pdf_url or (
            "https://www.google.com/search?q="
            + quote_plus(f"{f.company} {f.quarter} financial results pdf")
        )
        link.style = "Hyperlink"
        unit = fs.cell(r, 9, f.unit)
        unit.fill = entry_fill
        unit_dv.add(unit)
        fs.cell(r, 10, f'=COUNTIFS(Values!$L:$L,A{r},Values!$I:$I,"<>")')
        fs.cell(r, 11, f"=COUNTIF(Values!$L:$L,A{r})")
        fs.cell(r, 12, f.filing_id)
    fs.auto_filter.ref = f"A1:L{len(filings) + 1}"

    answer_dv = DataValidation(
        type="list", formula1=f"Lists!$D$1:$D${len(ANSWERS)}", allow_blank=True
    )
    vs.add_data_validation(answer_dv)
    r = 1
    for f in filings:
        statement = "Consolidated" if f.consolidated else "Standalone"
        for line in f.lines:
            r += 1
            unit_of = f"INDEX(Filings!$I:$I,MATCH(L{r},Filings!$A:$A,0))"
            values = (
                r - 1,
                f.symbol,
                statement,
                f.quarter,
                line.label,
                line.period,
                f'=IF(Q{r}="","not found",'
                f'IF(R{r}="money",Q{r}/VLOOKUP({unit_of},{factor_table},2,FALSE),Q{r}))',
                f'=IF(R{r}="money",{unit_of},IF(R{r}="per_share","₹ per share","as tagged"))',
                None,
                None,
                None,
                f.file,
                line.line,
                line.fact.element if line.fact else None,
                line.start.isoformat() if line.start else "",
                line.end.isoformat(),
                float(line.fact.value) if line.fact else None,
                line.kind,
                f.taxonomy,
            )
            for c, content in enumerate(values, start=1):
                if content is not None:
                    vs.cell(r, c, content)
            vs.cell(r, 7).number_format = "#,##0.00;[Red]-#,##0.00"
            for c in (9, 10, 11):
                vs.cell(r, c).fill = entry_fill
            answer_dv.add(vs.cell(r, 9))
    last = max(r, 2)
    vs.auto_filter.ref = f"A1:K{last}"
    for answer, colour in (("Yes", "C6EFCE"), ("No", "FFC7CE"), ("Not in PDF", "E7E6E6")):
        vs.conditional_formatting.add(
            f"A2:K{last}",
            FormulaRule(formula=[f'$I2="{answer}"'], fill=PatternFill("solid", bgColor=colour)),
        )
    wb.active = wb.sheetnames.index("Filings")
    wb.save(path)


@dataclass
class ImportSummary:
    confirmed: int = 0
    corrected: int = 0
    not_in_pdf: int = 0
    pending: int = 0
    problems: list[str] = field(default_factory=list)


def _plain(value: Decimal) -> str:
    return format(value.normalize(), "f")


def import_sheet(folder: Path = DEFAULT_DIR) -> ImportSummary:
    """Write manifest.yaml from the marked worksheet. Unmarked lines are left out (and counted)."""
    from openpyxl import load_workbook

    wb = load_workbook(folder / WORKSHEET)
    units = {
        row[0]: row[8]
        for row in wb["Filings"].iter_rows(min_row=2, values_only=True)
        if row and row[0]
    }
    summary = ImportSummary()
    entries: dict[str, dict[str, Any]] = {}
    parsed: dict[str, ParsedFiling] = {}
    for n, row in enumerate(wb["Values"].iter_rows(min_row=2, values_only=True), start=2):
        if not row or not row[11]:
            continue
        answer, shown_value, note = row[8], row[9], row[10]
        file, line, start, end = str(row[11]), str(row[12]), row[14], str(row[15])
        kind, taxonomy = row[17], str(row[18])
        where = f"Values row {n} ({row[1]} {row[4]}, {row[5]})"
        if not answer:
            summary.pending += 1
            continue
        if answer == "Not in PDF":
            summary.not_in_pdf += 1
            continue
        if file not in parsed:
            parsed[file] = parse_instance((folder / file).read_bytes())
        p_start = date.fromisoformat(str(start)) if start else None
        p_end = date.fromisoformat(end)
        fact = resolve(parsed[file], line, taxonomy, p_start, p_end)
        expect: dict[str, Any] = {"line": line}
        if p_start:
            expect["period_start"] = p_start
        expect["period_end"] = p_end
        if answer == "Yes":
            if fact is None:
                summary.problems.append(
                    f"{where}: Analekh has no value here; mark No and type the PDF's figure"
                )
                continue
            expect["value"] = _plain(fact.value)
            expect["checked"] = "confirmed"
            summary.confirmed += 1
        elif answer == "No":
            try:
                shown = Decimal(str(shown_value))
            except (InvalidOperation, ValueError):
                summary.problems.append(f"{where}: marked No without the PDF's figure")
                continue
            factor = UNITS.get(str(units.get(file)), 1) if kind == "money" else 1
            value = shown * factor
            if fact is not None and abs(fact.value - value) <= Decimal(factor) / 200:
                summary.problems.append(
                    f"{where}: marked No, but the PDF's figure agrees to two places; check again"
                )
            expect["value"] = _plain(value)
            expect["checked"] = "corrected"
            expect["analekh"] = _plain(fact.value) if fact else None
            summary.corrected += 1
        else:
            summary.problems.append(f"{where}: unknown answer {answer!r}")
            continue
        if note:
            expect["note"] = str(note)
        entry = entries.setdefault(file, {"file": file, "taxonomy": taxonomy, "expect": []})
        entry["expect"].append(expect)
    (folder / "manifest.yaml").write_text(
        "# Written by `analekh golden-import` from golden-worksheet.xlsx; edit the worksheet.\n"
        + yaml.safe_dump({"filings": list(entries.values())}, sort_keys=False, allow_unicode=True),
        encoding="utf-8",
    )
    return summary
